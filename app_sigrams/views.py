from datetime import timezone
from django.shortcuts import render, redirect, get_object_or_404
from django.db.models import Q
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.contrib import messages
from django.contrib.gis.geos import GEOSGeometry, GEOSException
from django.views.decorators.csrf import csrf_protect
from django.views.decorators.http import require_POST
from app_registros.models import UserProfile, Productor, MarcaSenal, Solicitud, Campo, TipoSenal, ImagenMarcaPredefinida
from app_registros.forms import ProductorForm, MarcaSenalForm, SolicitudForm
from django.http import JsonResponse
from django.views.generic import ListView, CreateView, UpdateView, DetailView
from django.urls import reverse_lazy, reverse
from django.conf import settings

from django.views.generic import ListView, CreateView, UpdateView, DetailView, DeleteView, TemplateView
from django.views import View
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from app_registros.models import Productor, Campo, MarcaSenal


from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from django.http import HttpResponse


import datetime
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from django.db.models import Sum, Count
from django.db.models.functions import TruncMonth

from reportlab.lib.pagesizes import landscape, A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from django_ratelimit.decorators import ratelimit

# Create your views here.

from django.http import JsonResponse
import json

def get_campos_por_productor(request, productor_id):
    """Obtener campos de un productor específico (para AJAX)"""
    try:
        productor = Productor.objects.get(id=productor_id)

        # Solo listar campos existentes. Ya NO se autocrea ninguno:
        # el campo se crea explícitamente cuando el usuario dibuja
        # su polígono en el formulario de Marca (ver crear_campo_poligono).
        campos = Campo.objects.filter(productor_id=productor_id)

        # Preparar datos para JSON, incluyendo el polígono (GeoJSON) si existe
        data = []
        for campo in campos:
            data.append({
                'id': campo.id,
                'nombre': campo.nombre,
                'distrito': campo.distrito or '',
                'departamento': campo.departamento or '',
                'poligono': json.loads(campo.poligono.geojson) if campo.poligono else None,
            })

        return JsonResponse(data, safe=False)

    except Productor.DoesNotExist:
        return JsonResponse([], safe=False)
    except Exception as e:
        print(f"Error en get_campos_por_productor: {str(e)}")
        return JsonResponse([{'error': str(e)}], safe=False, status=500)
    

@login_required
@require_POST
def crear_campo_poligono(request):
    """
    Crea un Campo a partir de un polígono dibujado en el mapa (Leaflet Draw).
    Espera POST con: productor_id, nombre, distrito, departamento,
    area_hectareas (opcional) y poligono_geojson (string GeoJSON del polígono).
    """
    try:
        productor_id = request.POST.get('productor_id')
        nombre = request.POST.get('nombre', '').strip()
        distrito = request.POST.get('distrito', '').strip()
        departamento = request.POST.get('departamento', '').strip()
        area_hectareas = request.POST.get('area_hectareas') or None
        poligono_geojson = request.POST.get('poligono_geojson')

        if not productor_id:
            return JsonResponse({'ok': False, 'error': 'Falta el productor.'}, status=400)
        if not nombre:
            return JsonResponse({'ok': False, 'error': 'El nombre del campo es obligatorio.'}, status=400)
        if not poligono_geojson:
            return JsonResponse({'ok': False, 'error': 'Debe dibujar el polígono del campo en el mapa.'}, status=400)

        productor = get_object_or_404(Productor, pk=productor_id)

        try:
            geom = GEOSGeometry(poligono_geojson)
        except (GEOSException, ValueError, TypeError):
            return JsonResponse({'ok': False, 'error': 'El polígono recibido no es válido.'}, status=400)

        if geom.geom_type != 'Polygon':
            return JsonResponse({'ok': False, 'error': 'La geometría debe ser un polígono.'}, status=400)

        if geom.srid is None:
            geom.srid = 4326

        centroide = geom.centroid

        campo = Campo.objects.create(
            nombre=nombre,
            productor=productor,
            distrito=distrito or "Sin especificar",
            departamento=departamento or "Sin especificar",
            area_hectareas=area_hectareas,
            poligono=geom,
            latitud=centroide.y,
            longitud=centroide.x,
        )

        return JsonResponse({
            'ok': True,
            'campo': {
                'id': campo.id,
                'nombre': campo.nombre,
                'distrito': campo.distrito or '',
                'departamento': campo.departamento or '',
                'poligono': json.loads(campo.poligono.geojson),
            }
        })

    except Exception as e:
        print(f"Error en crear_campo_poligono: {str(e)}")
        return JsonResponse({'ok': False, 'error': str(e)}, status=500)


@login_required
@require_POST
def actualizar_campo_poligono(request, campo_id):
    """
    Actualiza el polígono (y opcionalmente el nombre) de un Campo existente,
    a partir de una edición hecha en el mapa (Leaflet Draw - modo edición).
    """
    try:
        campo = get_object_or_404(Campo, pk=campo_id)

        poligono_geojson = request.POST.get('poligono_geojson')
        nombre = request.POST.get('nombre', '').strip()

        if not poligono_geojson:
            return JsonResponse({'ok': False, 'error': 'Falta el polígono editado.'}, status=400)

        try:
            geom = GEOSGeometry(poligono_geojson)
        except (GEOSException, ValueError, TypeError):
            return JsonResponse({'ok': False, 'error': 'El polígono recibido no es válido.'}, status=400)

        if geom.geom_type != 'Polygon':
            return JsonResponse({'ok': False, 'error': 'La geometría debe ser un polígono.'}, status=400)

        if geom.srid is None:
            geom.srid = 4326

        centroide = geom.centroid

        distrito = request.POST.get('distrito', '').strip()
        departamento = request.POST.get('departamento', '').strip()
        area_hectareas = request.POST.get('area_hectareas')

        campo.poligono = geom
        campo.latitud = centroide.y
        campo.longitud = centroide.x
        if nombre:
            campo.nombre = nombre
        if distrito:
            campo.distrito = distrito
        if departamento:
            campo.departamento = departamento
        if area_hectareas:
            campo.area_hectareas = area_hectareas
        campo.save()

        return JsonResponse({
            'ok': True,
            'campo': {
                'id': campo.id,
                'nombre': campo.nombre,
                'distrito': campo.distrito or '',
                'departamento': campo.departamento or '',
                'poligono': json.loads(campo.poligono.geojson),
            }
        })

    except Exception as e:
        print(f"Error en actualizar_campo_poligono: {str(e)}")
        return JsonResponse({'ok': False, 'error': str(e)}, status=500)
    
    
def get_imagenes_marcas(request):
    """Obtener imágenes predefinidas de marcas (para AJAX)"""
    try:
        imagenes = ImagenMarcaPredefinida.objects.filter(activa=True)
        data = []
        for imagen in imagenes:
            data.append({
                'id': imagen.id,
                'nombre': imagen.nombre,
                'tipo_marca': imagen.get_tipo_marca_display(),
                'imagen': imagen.imagen.url if imagen.imagen else '',
                'imagen_url': request.build_absolute_uri(imagen.imagen.url) if imagen.imagen else ''
            })
        return JsonResponse(data, safe=False)
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=400)


from django.http import JsonResponse
from app_registros.models import MarcaSenal, Productor
import json

def get_marcas_por_productor(request, productor_id):
    """Obtener marcas de un productor específico (para AJAX)"""
    try:
        print(f"DEBUG SERVIDOR: Solicitando marcas para productor_id={productor_id}")
        
        # Verificar que el productor existe
        if not Productor.objects.filter(id=productor_id).exists():
            return JsonResponse([], safe=False)
        
        # Obtener marcas del productor
        marcas = MarcaSenal.objects.filter(productor_id=productor_id).order_by('-fecha_inscripcion')
        
        # Preparar datos para JSON
        data = []
        for marca in marcas:
            data.append({
                'id': marca.id,
                'numero_orden': marca.numero_orden,
                'descripcion_marca': marca.descripcion_marca,
                'tipo_tramite': marca.get_tipo_tramite_display(),
                'estado': marca.get_estado_display(),
                'fecha_inscripcion': marca.fecha_inscripcion.strftime('%d/%m/%Y') if marca.fecha_inscripcion else 'No especificada',
            })
        
        print(f"DEBUG SERVIDOR: Enviando {len(data)} marcas")
        return JsonResponse(data, safe=False)
        
    except Exception as e:
        print(f"ERROR SERVIDOR en get_marcas_por_productor: {str(e)}")
        import traceback
        traceback.print_exc()
        return JsonResponse([{'error': str(e)}], safe=False, status=500)


@login_required
def detalle_marca_ajax(request, marca_id):
    """Devuelve el detalle de una marca en JSON (para mostrarla como 'actual' en Solicitud)."""
    try:
        marca = MarcaSenal.objects.get(pk=marca_id)
        return JsonResponse({
            'ok': True,
            'numero_orden': marca.numero_orden,
            'descripcion_marca': marca.descripcion_marca,
            'descripcion_senal': marca.descripcion_senal or '',
            'tipo_marca': marca.tipo_marca or '',
            'tipo_marca_display': marca.get_tipo_marca_display() if marca.tipo_marca else '',
            'senales_orejeras': marca.senales_orejeras or {},
            'estado': marca.get_estado_display(),
            'imagen_url': marca.imagen_marca.url if marca.imagen_marca else None,
        })
    except MarcaSenal.DoesNotExist:
        return JsonResponse({'ok': False, 'error': 'Marca no encontrada'}, status=404)


@login_required
def verificar_marcas_productor_ajax(request, productor_id):
    """Devuelve cuántas marcas tiene un productor (para avisar en transferencias)."""
    try:
        productor = Productor.objects.get(pk=productor_id)
        cantidad = productor.marcas_senales.count()
        return JsonResponse({
            'ok': True,
            'tiene_marcas': cantidad > 0,
            'cantidad': cantidad,
            'nombre': str(productor),
        })
    except Productor.DoesNotExist:
        return JsonResponse({'ok': False, 'error': 'Productor no encontrado'}, status=404)


# app_sigrams/views.py - Actualizar función home
from django.contrib.auth.decorators import login_required
from django.shortcuts import render
from django.utils import timezone
from django.db.models import Count, Sum
import datetime
import json


@login_required
def home(request):
    """Vista principal del dashboard"""
    from app_registros.models import Productor, MarcaSenal, Solicitud, Campo

    hoy = timezone.now()

    # ========== ESTADÍSTICAS BÁSICAS ==========
    total_productores = Productor.objects.count()
    total_marcas = MarcaSenal.objects.filter(estado='VIGENTE').count()
    solicitudes_pendientes = Solicitud.objects.filter(estado='PENDIENTE').count()
    total_campos = Campo.objects.count()

    # ========== INGRESOS ==========
    # Primer día del mes actual
    primer_dia_mes = hoy.replace(day=1, hour=0, minute=0, second=0, microsecond=0)

    # Ingresos mes actual
    ingresos_mes = (
        MarcaSenal.objects.filter(
            fecha_creacion__gte=primer_dia_mes,
            valor_sellado__isnull=False
        ).aggregate(total=Sum('valor_sellado'))['total']
        or 0
    )

    # Ingresos mes anterior
    primer_dia_mes_anterior = (primer_dia_mes - datetime.timedelta(days=1)).replace(day=1)
    ultimo_dia_mes_anterior = primer_dia_mes - datetime.timedelta(days=1)

    ingresos_mes_anterior = (
        MarcaSenal.objects.filter(
            fecha_creacion__gte=primer_dia_mes_anterior,
            fecha_creacion__lte=ultimo_dia_mes_anterior,
            valor_sellado__isnull=False
        ).aggregate(total=Sum('valor_sellado'))['total']
        or 0
    )

    # Variación porcentual
    if ingresos_mes_anterior > 0:
        variacion_ingresos = ((ingresos_mes - ingresos_mes_anterior) / ingresos_mes_anterior) * 100
    else:
        variacion_ingresos = 100 if ingresos_mes > 0 else 0

    # ========== TRÁMITES ==========
    tramites_mes = Solicitud.objects.filter(
        fecha_solicitud__gte=primer_dia_mes
    ).count()

    # ========== DISTRIBUCIONES ==========
    productores_por_estado = list(
        Productor.objects.values('estado')
        .annotate(total=Count('id'))
        .order_by('estado')
    )

    marcas_por_tipo = list(
        MarcaSenal.objects.values('tipo_tramite')
        .annotate(total=Count('id'))
        .order_by('tipo_tramite')
    )

    solicitudes_por_estado = list(
        Solicitud.objects.values('estado')
        .annotate(total=Count('id'))
        .order_by('estado')
    )

    # ========== DATOS RECIENTES ==========
    ultimas_solicitudes = (
        Solicitud.objects.select_related('productor')
        .order_by('-fecha_solicitud')[:10]
    )

    fecha_limite = hoy - datetime.timedelta(days=7)
    productores_recientes = (
        Productor.objects.filter(fecha_registro__gte=fecha_limite)
        .order_by('-fecha_registro')[:5]
    )

    fecha_vencimiento = hoy + datetime.timedelta(days=30)
    marcas_por_vencer = (
        MarcaSenal.objects.filter(
            fecha_vencimiento__gte=hoy,
            fecha_vencimiento__lte=fecha_vencimiento,
            estado='VIGENTE'
        )
        .select_related('productor')
        .order_by('fecha_vencimiento')[:5]
    )

    marcas_recientes = (
        MarcaSenal.objects.select_related('productor')
        .order_by('-fecha_inscripcion')[:5]
    )

    # ========== ACTIVIDAD MENSUAL (ÚLTIMOS 6 MESES) ==========
    meses_actividad = []

    for i in range(6):
        fecha = hoy - datetime.timedelta(days=30 * i)
        mes_str = fecha.strftime('%Y-%m')
        nombre_mes = fecha.strftime('%b')

        solicitudes_mes = Solicitud.objects.filter(
            fecha_solicitud__year=fecha.year,
            fecha_solicitud__month=fecha.month
        ).count()

        marcas_mes = MarcaSenal.objects.filter(
            fecha_inscripcion__year=fecha.year,
            fecha_inscripcion__month=fecha.month
        ).count()

        ingresos_mes_valor = (
            MarcaSenal.objects.filter(
                fecha_creacion__year=fecha.year,
                fecha_creacion__month=fecha.month,
                valor_sellado__isnull=False
            ).aggregate(total=Sum('valor_sellado'))['total']
            or 0
        )

        meses_actividad.append({
            'mes': mes_str,
            'nombre': nombre_mes,
            'solicitudes': solicitudes_mes,
            'marcas': marcas_mes,
            'ingresos': float(ingresos_mes_valor)
        })

    meses_actividad.reverse()

    # ========== TOP PRODUCTORES ==========
    top_productores = (
        Productor.objects.annotate(total_marcas=Count('marcas_senales'))
        .order_by('-total_marcas')[:5]
    )

    # ========== CONTEXT ==========
    context = {
        # Básicos
        'total_productores': total_productores,
        'total_marcas': total_marcas,
        'solicitudes_pendientes': solicitudes_pendientes,
        'total_campos': total_campos,

        # Finanzas
        'ingresos_mes': ingresos_mes,
        'ingresos_mes_anterior': ingresos_mes_anterior,
        'variacion_ingresos': variacion_ingresos,
        'tramites_mes': tramites_mes,

        # Gráficos
        'productores_por_estado': json.dumps(productores_por_estado),
        'marcas_por_tipo': json.dumps(marcas_por_tipo),
        'solicitudes_por_estado': json.dumps(solicitudes_por_estado),
        'meses_actividad': json.dumps(meses_actividad),

        # Listados
        'ultimas_solicitudes': ultimas_solicitudes,
        'productores_recientes': productores_recientes,
        'marcas_por_vencer': marcas_por_vencer,
        'marcas_recientes': marcas_recientes,
        'top_productores': top_productores,

        # Badges / listas
        'productores_estados_list': productores_por_estado,
        'marcas_tipos_list': marcas_por_tipo,
    }

    return render(request, 'app_sigrams/index.html', context)


from django.db.models.functions import TruncMonth
from django.db.models import Count, Sum
from django.utils import timezone
from datetime import timedelta

@login_required
def dashboard_admin(request):
    user_profile = UserProfile.objects.get(user=request.user)
    if user_profile.rol != 'admin':
        messages.error(request, 'No tiene permisos para acceder.')
        return redirect('home')

    # Usuarios
    total_usuarios = User.objects.count()
    admins = UserProfile.objects.filter(rol='admin').count()
    empleados = UserProfile.objects.filter(rol='empleado').count()

    # Solicitudes
    total_solicitudes = Solicitud.objects.count()
    solicitudes_pendientes = Solicitud.objects.filter(estado='PENDIENTE').count()
    solicitudes_en_revision = Solicitud.objects.filter(estado='EN_REVISION').count()
    solicitudes_aprobadas = Solicitud.objects.filter(estado='APROBADO').count()
    solicitudes_rechazadas = Solicitud.objects.filter(estado='RECHAZADO').count()

    # Productores y Marcas
    total_productores = Productor.objects.count()
    total_marcas = MarcaSenal.objects.count()

    # Ingresos (últimos 12 meses)
    hoy = timezone.now()
    doce_meses = hoy - timedelta(days=365)
    ingresos_por_mes = (
        MarcaSenal.objects
        .filter(fecha_creacion__gte=doce_meses, valor_sellado__isnull=False)
        .annotate(mes=TruncMonth('fecha_creacion'))
        .values('mes')
        .annotate(total=Sum('valor_sellado'))
        .order_by('mes')
    )

    # Logs recientes (últimos 20)
    logs_recientes = ChangeLog.objects.select_related('user').order_by('-timestamp')[:20]

    # Últimas solicitudes
    ultimas_solicitudes = Solicitud.objects.select_related('productor', 'solicitante').order_by('-fecha_solicitud')[:10]

    context = {
        'total_usuarios': total_usuarios,
        'admins': admins,
        'empleados': empleados,
        'total_solicitudes': total_solicitudes,
        'solicitudes_pendientes': solicitudes_pendientes,
        'solicitudes_en_revision': solicitudes_en_revision,
        'solicitudes_aprobadas': solicitudes_aprobadas,
        'solicitudes_rechazadas': solicitudes_rechazadas,
        'total_productores': total_productores,
        'total_marcas': total_marcas,
        'ingresos_por_mes': list(ingresos_por_mes),
        'logs_recientes': logs_recientes,
        'ultimas_solicitudes': ultimas_solicitudes,
    }
    return render(request, 'app_sigrams/admin/dashboard.html', context)



# LOGIN
@ratelimit(key='ip', rate='60/h', method='POST', block=False)
@ratelimit(key='post:username', rate='10/h', method='POST', block=False)
def login_view(request):
    if request.user.is_authenticated:
        return redirect('home')

    if request.method == 'POST':
        if getattr(request, 'limited', False):
            messages.error(request, 'Demasiados intentos de inicio de sesión. Esperá un rato antes de volver a probar.')
            return render(request, 'app_sigrams/login.html')

        # Verificar reCAPTCHA
        recaptcha_response = request.POST.get('g-recaptcha-response')
        recaptcha_error = None

        if not recaptcha_response:
            recaptcha_error = 'Por favor completá el reCAPTCHA.'
        else:
            import urllib.request
            import urllib.parse
            import json as _json
            private_key = settings.RECAPTCHA_PRIVATE_KEY
            data = urllib.parse.urlencode({
                'secret': private_key,
                'response': recaptcha_response,
                'remoteip': request.META.get('REMOTE_ADDR'),
            }).encode('utf-8')
            try:
                req = urllib.request.urlopen(
                    'https://www.google.com/recaptcha/api/siteverify',
                    data=data,
                    timeout=5
                )
                result = _json.loads(req.read().decode('utf-8'))
                if not result.get('success'):
                    recaptcha_error = 'Verificación reCAPTCHA fallida. Intentá de nuevo.'
            except Exception:
                recaptcha_error = 'No se pudo verificar el reCAPTCHA. Intentá de nuevo.'

        if recaptcha_error:
            return render(request, 'app_sigrams/login.html', {
                'recaptcha_error': recaptcha_error
            })

        username = request.POST.get('username', '')
        password = request.POST.get('password', '')
        user = authenticate(request, username=username, password=password)

        if user is not None:
            login(request, user)
            messages.success(request, f'¡Bienvenido/a {user.username}!')
            return redirect('home')
        else:
            # Solo se informa el estado de la cuenta si la contraseña es correcta;
            # en cualquier otro caso, el mismo mensaje genérico.
            posible_user = User.objects.filter(username=username).first()
            if (
                posible_user
                and not posible_user.is_active
                and posible_user.check_password(password)
            ):
                profile = UserProfile.objects.filter(user=posible_user).first()
                estado = profile.estado_cuenta if profile else None

                if estado == 'PENDIENTE':
                    messages.warning(
                        request,
                        'Tu cuenta todavía está pendiente de aprobación por un administrador.'
                    )
                elif estado == 'RECHAZADA':
                    messages.error(
                        request,
                        'Tu solicitud de cuenta fue rechazada. Contactá a un administrador.'
                    )
                elif estado == 'APROBADA':
                    messages.warning(
                        request,
                        'Tu cuenta fue aprobada pero todavía no fue activada. '
                        'Pedile el código de activación al administrador.'
                    )
                else:
                    messages.error(request, 'Usuario o contraseña incorrectos.')
            else:
                messages.error(request, 'Usuario o contraseña incorrectos.')

    return render(request, 'app_sigrams/login.html')


# LOGOUT
@login_required
def logout_view(request):
    logout(request)
    messages.info(request, 'Has cerrado sesión correctamente.')
    return redirect('login')


# REGISTRO
from django.shortcuts import render, redirect
from django.contrib import messages
from django.contrib.auth.models import User

from app_sigrams.forms import CustomUserCreationForm

@ratelimit(key='ip', rate='10/h', method='POST', block=False)
@ratelimit(key='post:email', rate='3/h', method='POST', block=False)
def register_view(request):
    if request.user.is_authenticated:
        return redirect('home')

    if request.method == 'POST':
        if getattr(request, 'limited', False):
            messages.error(request, 'Demasiados intentos de registro. Esperá un rato antes de volver a probar.')
            return render(request, 'app_sigrams/register.html', {'form': CustomUserCreationForm()})
        # Verificar reCAPTCHA
        recaptcha_response = request.POST.get('g-recaptcha-response')
        recaptcha_error = None

        if not recaptcha_response:
            recaptcha_error = 'Por favor completá el reCAPTCHA.'
        else:
            import urllib.request
            import urllib.parse
            import json as _json
            data = urllib.parse.urlencode({
                'secret': settings.RECAPTCHA_PRIVATE_KEY,
                'response': recaptcha_response,
                'remoteip': request.META.get('REMOTE_ADDR'),
            }).encode('utf-8')
            try:
                req = urllib.request.urlopen(
                    'https://www.google.com/recaptcha/api/siteverify',
                    data=data,
                    timeout=5
                )
                result = _json.loads(req.read().decode('utf-8'))
                if not result.get('success'):
                    recaptcha_error = 'Verificación reCAPTCHA fallida. Intentá de nuevo.'
            except Exception:
                recaptcha_error = 'No se pudo verificar el reCAPTCHA. Intentá de nuevo.'

        if recaptcha_error:
            form = CustomUserCreationForm(request.POST)
            return render(request, 'app_sigrams/register.html', {
                'form': form,
                'recaptcha_error': recaptcha_error
            })

        form = CustomUserCreationForm(request.POST)

        if form.is_valid():
            user = form.save(commit=False)
            user.is_active = False  # queda inactivo hasta que un admin lo apruebe
            user.save()

            # Crear o obtener perfil (evita duplicados), en estado pendiente
            profile, created = UserProfile.objects.get_or_create(
                user=user,
                defaults={'rol': 'empleado', 'estado_cuenta': 'PENDIENTE'}
            )
            if not created:
                profile.estado_cuenta = 'PENDIENTE'
                profile.save()

            # Notificar a los administradores (in-app + email)
            from app_notificaciones.models import Notificacion
            admins = User.objects.filter(userprofile__rol='admin', is_active=True)
            for admin_user in admins:
                Notificacion.crear_notificacion(
                    usuario=admin_user,
                    titulo='Nueva solicitud de cuenta',
                    mensaje=f'{user.get_full_name() or user.username} solicitó una cuenta y espera aprobación.',
                    tipo='alerta',
                    contenido_relacionado=profile,
                    url='/cuentas-pendientes/'
                )

            admin_emails = [a.email for a in admins if a.email]
            if admin_emails:
                try:
                    from django.core.mail import send_mail
                    send_mail(
                        subject='SIGRAMS — Nueva solicitud de cuenta pendiente',
                        message=(
                            f'{user.get_full_name() or user.username} ({user.email}) '
                            f'solicitó una cuenta en SIGRAMS y espera aprobación.\n\n'
                            f'Ingresá al sistema para revisar y aprobar la solicitud.'
                        ),
                        from_email=settings.DEFAULT_FROM_EMAIL,
                        recipient_list=admin_emails,
                        fail_silently=True,
                    )
                except Exception:
                    pass

            messages.success(
                request,
                'Tu cuenta fue creada y está pendiente de aprobación por un administrador. '
                'Te van a contactar con un código de activación cuando sea aprobada.'
            )
            return redirect('login')

        else:
            # Mostrar errores campo por campo como mensajes
            for field, errors in form.errors.items():
                label = (
                    form.fields[field].label
                    if field in form.fields
                    else field
                )
                for error in errors:
                    messages.error(request, f"{label}: {error}")

    else:
        form = CustomUserCreationForm()

    return render(
        request,
        'app_sigrams/register.html',
        {'form': form}
    )



class EsEmpleadoOMas(UserPassesTestMixin):
    def test_func(self):
        user_profile = UserProfile.objects.get(user=self.request.user)
        return user_profile.rol in ['admin', 'empleado']



# ============================================================================
# VISTAS PARA PRODUCTORES
# ============================================================================
from django.shortcuts import render, redirect, get_object_or_404
from django.core.paginator import Paginator

@login_required
def lista_productores(request):
    """Lista todos los productores con filtros"""
    productores = Productor.objects.all()

    # Filtros — coinciden con los names del formulario en lista.html
    nombre = request.GET.get('nombre', '')
    dni = request.GET.get('dni', '')
    estado = request.GET.get('estado', '')
    localidad = request.GET.get('localidad', '')

    if nombre:
        productores = productores.filter(
            Q(nombre__icontains=nombre) |
            Q(apellido__icontains=nombre)
        )

    if dni:
        productores = productores.filter(dni__icontains=dni)

    if estado:
        productores = productores.filter(estado=estado)

    if localidad:
        productores = productores.filter(localidad__icontains=localidad)

    total_productores = productores.count()

    paginator = Paginator(productores, 10)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    context = {
        'productores': page_obj,
        'page_obj': page_obj,
        'is_paginated': page_obj.has_other_pages(),
        'total_productores': total_productores,
        'estados': Productor.ESTADO_CHOICES,
    }
    return render(request, 'app_sigrams/productores/lista.html', context)

@login_required
def detalle_productor(request, pk):
    """Detalle de un productor específico"""
    productor = get_object_or_404(Productor, pk=pk)
    marcas = productor.marcas_senales.all()
    solicitudes = productor.solicitudes.all()
    campos = productor.campos.all()

    # Historial: cambios sobre este productor, sus marcas, sus solicitudes,
    # y transferencias donde participó (como origen o como destino).
    ids_marcas = list(marcas.values_list('id', flat=True))
    ids_solicitudes = list(solicitudes.values_list('id', flat=True))
    ids_solicitudes_recibidas = list(
        Solicitud.objects.filter(productor_destino=productor).values_list('id', flat=True)
    )

    historial = ChangeLog.objects.filter(
        Q(modelo='Productor', objeto_id=str(productor.id)) |
        Q(modelo='MarcaSenal', objeto_id__in=[str(i) for i in ids_marcas]) |
        Q(modelo='Solicitud', objeto_id__in=[str(i) for i in ids_solicitudes + ids_solicitudes_recibidas]) |
        Q(accion='TRANSFERENCIA_POR_SOLICITUD', snapshot__anterior__productor_id=productor.id) |
        Q(accion='TRANSFERENCIA_POR_SOLICITUD', snapshot__nuevo__productor_id=productor.id)
    ).distinct().order_by('-timestamp')[:50]

    # Marcas que este productor tuvo y transfirió a otro (para mostrarlas
    # en su detalle aunque ya no le pertenezcan).
    marcas_transferidas = ChangeLog.objects.filter(
        modelo='MarcaSenal',
        accion='TRANSFERENCIA_POR_SOLICITUD',
        snapshot__anterior__productor_id=productor.id,
    ).order_by('-timestamp')

    context = {
        'productor': productor,
        'marcas': marcas,
        'solicitudes': solicitudes,
        'campos': campos,
        'historial': historial,
        'marcas_transferidas': marcas_transferidas,
    }
    return render(request, 'app_sigrams/productores/detalle.html', context)

@login_required
def nuevo_productor(request):
    """Crear nuevo productor"""
    try:
        if request.method == 'POST':
            form = ProductorForm(request.POST)
            if form.is_valid():
                productor = form.save()
                messages.success(request, f'Productor {productor.nombre_completo} creado exitosamente.')
                return redirect('detalle_productor', pk=productor.pk)
            else:
                # Obtener todos los errores para mostrarlos
                error_list = []
                for field, errors in form.errors.items():
                    for error in errors:
                        if field != '__all__':
                            field_label = form.fields[field].label or field
                            error_list.append(f"{field_label}: {error}")
                        else:
                            error_list.append(str(error))
                
                if error_list:
                    messages.error(request, "Errores en el formulario:")
                    for error in error_list:
                        messages.error(request, f"- {error}")
                else:
                    messages.error(request, 'Por favor, corrija los errores en el formulario.')
                
                # Mantener los datos ingresados
                context = {'form': form, 'titulo': 'Nuevo Productor'}
                return render(request, 'app_sigrams/productores/form.html', context)
        else:
            form = ProductorForm()
        
        context = {'form': form, 'titulo': 'Nuevo Productor'}
        return render(request, 'app_sigrams/productores/form.html', context)
    
    except Exception as e:
        messages.error(request, f'Error al crear el productor: {str(e)}')
        return redirect('lista_productores')



@login_required
def editar_productor(request, pk):
    """Editar productor existente"""
    try:
        productor = get_object_or_404(Productor, pk=pk)
        
        if request.method == 'POST':
            form = ProductorForm(request.POST, instance=productor)
            if form.is_valid():
                productor = form.save()
                messages.success(request, f'Productor {productor.nombre_completo} actualizado exitosamente.')
                return redirect('detalle_productor', pk=productor.pk)
            else:
                messages.error(request, 'Por favor, corrija los errores en el formulario.')
        else:
            form = ProductorForm(instance=productor)
        
        context = {'form': form, 'titulo': 'Editar Productor', 'productor': productor}
        return render(request, 'app_sigrams/productores/form.html', context)
    
    except Exception as e:
        messages.error(request, f'Error al editar el productor: {str(e)}')
        return redirect('lista_productores')



@login_required
def eliminar_productor(request, pk):
    """Eliminar productor (solo administradores)."""
    default_rol = 'admin' if request.user.is_superuser else 'empleado'
    user_profile, _ = UserProfile.objects.get_or_create(
        user=request.user,
        defaults={'rol': default_rol}
    )
    if user_profile.rol != 'admin':
        messages.error(request, 'No tiene permisos para eliminar productores.')
        return redirect('detalle_productor', pk=pk)

    productor = get_object_or_404(Productor, pk=pk)

    if request.method == 'POST':
        nombre = productor.nombre_completo
        dni = productor.dni

        # Guardar constancia detallada de todo lo que se elimina en cascada
        detalle_marcas = [
            {
                'numero_orden': m.numero_orden,
                'estado': m.get_estado_display(),
                'tipo_tramite': m.get_tipo_tramite_display(),
                'descripcion_marca': m.descripcion_marca[:150],
                'descripcion_senal': (m.descripcion_senal or '')[:150],
                'total_ganado': m.total_ganado,
            }
            for m in productor.marcas_senales.all()
        ]
        detalle_campos = [
            {
                'nombre': camp.nombre,
                'distrito': camp.distrito,
                'departamento': camp.departamento,
                'area_hectareas': str(camp.area_hectareas) if camp.area_hectareas else '',
            }
            for camp in productor.campos.all()
        ]
        detalle_solicitudes = [
            {
                'id': s.id,
                'tipo_tramite': s.get_tipo_tramite_display(),
                'estado': s.get_estado_display(),
                'fecha': s.fecha_solicitud.strftime('%d/%m/%Y'),
            }
            for s in productor.solicitudes.all()
        ]

        ChangeLog.objects.create(
            user=request.user,
            modelo='Productor',
            objeto_id=str(productor.pk),
            accion='ELIMINACION',
            snapshot={
                'nombre': nombre,
                'dni': dni,
                'cuit': productor.cuit or '',
                'localidad': productor.localidad or '',
                'departamento': productor.departamento or '',
                'estado': productor.get_estado_display(),
                'marcas': detalle_marcas,
                'campos': detalle_campos,
                'solicitudes': detalle_solicitudes,
            }
        )

        productor.delete()
        messages.success(request, f'Productor {nombre} eliminado exitosamente.')
        return redirect('lista_productores')

    context = {'productor': productor}
    return render(request, 'app_sigrams/productores/confirmar_eliminar.html', context)

# Vistas para Productores
class ListaProductoresView(LoginRequiredMixin, ListView):
    model = Productor
    template_name = 'productores/lista.html'
    context_object_name = 'productores'
    paginate_by = 20
    
    def get_queryset(self):
        queryset = super().get_queryset()
        # Filtros
        nombre = self.request.GET.get('nombre')
        dni = self.request.GET.get('dni')
        localidad = self.request.GET.get('localidad')
        estado = self.request.GET.get('estado')
        
        if nombre:
            queryset = queryset.filter(Q(nombre__icontains=nombre) | Q(apellido__icontains=nombre))
        if dni:
            queryset = queryset.filter(dni__icontains=dni)
        if localidad:
            queryset = queryset.filter(localidad__icontains=localidad)
        if estado:
            queryset = queryset.filter(estado=estado)
        
        return queryset.order_by('apellido', 'nombre')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['total_productores'] = Productor.objects.count()
        context['estados'] = Productor.ESTADO_CHOICES
        return context

class NuevoProductorView(LoginRequiredMixin, EsEmpleadoOMas, CreateView):
    model = Productor
    template_name = 'productores/form.html'
    fields = ['nombre', 'apellido', 'dni', 'cuit', 'domicilio', 'distrito', 
              'localidad', 'telefono', 'email', 'latitud', 'longitud', 
              'estado', 'observaciones']
    success_url = reverse_lazy('lista_productores')
    
    def form_valid(self, form):
        form.instance.usuario_registro = self.request.user
        return super().form_valid(form)

class EditarProductorView(LoginRequiredMixin, EsEmpleadoOMas, UpdateView):
    model = Productor
    template_name = 'productores/form.html'
    fields = ['nombre', 'apellido', 'dni', 'cuit', 'domicilio', 'distrito', 
              'localidad', 'telefono', 'email', 'latitud', 'longitud', 
              'estado', 'observaciones']
    success_url = reverse_lazy('lista_productores')

class DetalleProductorView(LoginRequiredMixin, DetailView):
    model = Productor
    template_name = 'productores/detalle.html'
    context_object_name = 'productor'
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['campos'] = self.object.campos.all()
        context['marcas'] = self.object.marcas_senales.all()
        context['solicitudes'] = self.object.solicitudes.all()
        return context

class EliminarProductorView(LoginRequiredMixin, UserPassesTestMixin, DeleteView):
    model = Productor
    template_name = 'productores/confirmar_eliminar.html'
    success_url = reverse_lazy('lista_productores')
    
    def test_func(self):
        user_profile = UserProfile.objects.get(user=self.request.user)
        return user_profile.rol == 'admin'


class ReporteProductoresPDFView(LoginRequiredMixin, View):
    def get(self, request):
        # Crear respuesta HTTP con contenido PDF
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="reporte_productores.pdf"'
        
        # Crear documento PDF
        doc = SimpleDocTemplate(response, pagesize=letter)
        elements = []
        
        # Estilos
        styles = getSampleStyleSheet()
        
        # Título
        title = Paragraph("Reporte de Productores - SIGRAMS", styles['Title'])
        elements.append(title)
        
        # Fecha
        fecha = Paragraph(f"Generado: {timezone.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal'])
        elements.append(fecha)
        
        elements.append(Paragraph("<br/><br/>", styles['Normal']))
        
        # Datos
        productores = Productor.objects.all().order_by('apellido', 'nombre')
        
        # Tabla
        data = [['Nombre', 'DNI', 'Localidad', 'Estado', 'Fecha Registro']]
        
        for p in productores:
            data.append([
                f"{p.apellido}, {p.nombre}",
                p.dni,
                p.localidad,
                p.get_estado_display(),
                p.fecha_registro.strftime('%d/%m/%Y')
            ])
        
        table = Table(data)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0A2E5A')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ]))
        
        elements.append(table)
        
        # Generar PDF
        doc.build(elements)
        
        return response

# ============================================================================
# VISTAS PARA MARCAS Y SEÑALES
# ============================================================================

# Vistas para Marcas y Señales
@login_required
def lista_marcas(request):
    """Lista todas las marcas y señales"""
    marcas = MarcaSenal.objects.select_related('productor', 'campo').all()
    
    # Filtros
    query = request.GET.get('q', '')
    estado = request.GET.get('estado', '')
    tipo_tramite = request.GET.get('tipo_tramite', '')
    
    if query:
        marcas = marcas.filter(
            Q(productor__nombre__icontains=query) | 
            Q(productor__apellido__icontains=query) |
            Q(numero_orden__icontains=query)
        )
    
    if estado:
        marcas = marcas.filter(estado=estado)
    
    if tipo_tramite:
        marcas = marcas.filter(tipo_tramite=tipo_tramite)
    
    context = {
        'marcas': marcas,
        'query': query,
        'estado_filtro': estado,
        'tipo_tramite_filtro': tipo_tramite,
    }
    return render(request, 'app_sigrams/marcas/lista.html', context)



@login_required
def detalle_marca(request, pk):
    """Detalle de una marca específica"""
    marca = get_object_or_404(MarcaSenal.objects.select_related('productor', 'campo', 'tipo_senal'), pk=pk)
    
    context = {
        'marca': marca,
    }
    return render(request, 'app_sigrams/marcas/detalle.html', context)



@login_required
def nueva_marca(request):
    """Crear nueva marca y señal"""

    if request.method == 'POST':
        form = MarcaSenalForm(request.POST, request.FILES)

        if form.is_valid():
            marca = form.save(commit=False)

            imagenes_predefinidas = form.cleaned_data.get('imagenes_predefinidas')

            # Si no subieron imagen manual pero eligieron predefinidas
            if imagenes_predefinidas and not request.FILES.get('imagen_marca'):
                primera_imagen = imagenes_predefinidas.first()

                if primera_imagen and primera_imagen.imagen:
                    contenido = primera_imagen.imagen.read()
                    nombre_archivo = os.path.basename(primera_imagen.imagen.name)

                    marca.imagen_marca.save(
                        f"predef_{primera_imagen.id}_{nombre_archivo}",
                        ContentFile(contenido),
                        save=False
                    )

            # 🔥 Aquí se genera automáticamente numero_orden
            marca.save()

            form.save_m2m()

            messages.success(
                request,
                f'Marca #{marca.numero_orden} creada exitosamente.'
            )

            return redirect('detalle_marca', pk=marca.pk)

        messages.error(request, 'Por favor, corrija los errores en el formulario.')

    else:
        form = MarcaSenalForm()

    return render(request, 'app_sigrams/marcas/form.html', {
        'form': form,
        'titulo': 'Nueva Marca y Señal',
        'imagenes_predefinidas': ImagenMarcaPredefinida.objects.filter(activa=True)
    })




@login_required
def editar_marca(request, pk):
    """Editar marca existente"""

    marca = get_object_or_404(MarcaSenal, pk=pk)

    if request.method == 'POST':
        form = MarcaSenalForm(request.POST, request.FILES, instance=marca)

        if form.is_valid():
            marca_temp = form.save(commit=False)

            imagenes_predefinidas = form.cleaned_data.get('imagenes_predefinidas')

            if (
                imagenes_predefinidas
                and not request.FILES.get('imagen_marca')
                and not marca.imagen_marca
            ):
                primera_imagen = imagenes_predefinidas.first()

                if primera_imagen and primera_imagen.imagen:
                    contenido = primera_imagen.imagen.read()
                    nombre_archivo = os.path.basename(primera_imagen.imagen.name)

                    marca_temp.imagen_marca.save(
                        f"predef_{primera_imagen.id}_{nombre_archivo}",
                        ContentFile(contenido),
                        save=False
                    )

            # 🔥 No modifica numero_orden porque ya existe
            marca_temp.save()

            form.save_m2m()

            messages.success(
                request,
                f'Marca #{marca_temp.numero_orden} actualizada exitosamente.'
            )

            return redirect('detalle_marca', pk=marca_temp.pk)

        messages.error(request, 'Por favor, corrija los errores en el formulario.')

    else:
        form = MarcaSenalForm(instance=marca)

    return render(request, 'app_sigrams/marcas/form.html', {
        'form': form,
        'titulo': f'Editar Marca #{marca.numero_orden}',
        'imagenes_predefinidas': ImagenMarcaPredefinida.objects.filter(activa=True),
        'marca': marca
    })



class ListaMarcasView(ListView):
    model = MarcaSenal
    template_name = 'app_sigrams/marcas/lista.html'
    context_object_name = 'marcas'
    paginate_by = 10

    def get_queryset(self):
        queryset = MarcaSenal.objects.select_related('productor').all()

        query = self.request.GET.get('q', '')
        estado = self.request.GET.get('estado', '')
        tipo_tramite = self.request.GET.get('tipo_tramite', '')

        if query:
            queryset = queryset.filter(
                Q(productor__nombre__icontains=query) |
                Q(productor__apellido__icontains=query) |
                Q(numero_orden__icontains=query)
            )

        if estado:
            queryset = queryset.filter(estado=estado)

        if tipo_tramite:
            queryset = queryset.filter(tipo_tramite=tipo_tramite)

        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['query'] = self.request.GET.get('q', '')
        context['estado_filtro'] = self.request.GET.get('estado', '')
        context['tipo_tramite_filtro'] = self.request.GET.get('tipo_tramite', '')
        return context

class NuevaMarcaView(CreateView):
    model = MarcaSenal
    form_class = MarcaSenalForm
    template_name = 'app_sigrams/marcas/form.html'
    success_url = reverse_lazy('lista_marcas')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['titulo'] = 'Nueva Marca y Señal'
        context['imagenes_predefinidas'] = ImagenMarcaPredefinida.objects.filter(activa=True)
        return context

    def form_valid(self, form):
        marca = form.save(commit=False)
        # Guardar señales orejeras desde el campo oculto del formulario
        senales_json = self.request.POST.get('senales_orejeras_json', '{}')
        try:
            import json
            marca.senales_orejeras = json.loads(senales_json)
        except Exception:
            marca.senales_orejeras = {}
        marca.save()
        form.save_m2m()
        messages.success(self.request, f'Marca #{marca.numero_orden} creada exitosamente.')
        return redirect(self.success_url)

    def form_invalid(self, form):
        messages.error(self.request, 'Por favor, corrija los errores en el formulario.')
        return super().form_invalid(form)

class EditarMarcaView(UpdateView):
    model = MarcaSenal
    form_class = MarcaSenalForm
    template_name = 'app_sigrams/marcas/form.html'
    success_url = reverse_lazy('lista_marcas')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['titulo'] = f'Editar Marca y Señal #{self.object.numero_orden}'
        context['imagenes_predefinidas'] = ImagenMarcaPredefinida.objects.filter(activa=True)
        # Pasar las señales guardadas para pre-cargar la herramienta
        context['senales_guardadas'] = json.dumps(
            self.object.senales_orejeras or {}
        )
        return context

    def form_valid(self, form):
        marca = form.save(commit=False)
        senales_json = self.request.POST.get('senales_orejeras_json', '{}')
        try:
            marca.senales_orejeras = json.loads(senales_json)
        except Exception:
            marca.senales_orejeras = {}
        marca.save()
        form.save_m2m()
        messages.success(self.request, f'Marca #{marca.numero_orden} actualizada exitosamente.')
        return redirect(self.success_url)

    def form_invalid(self, form):
        messages.error(self.request, 'Por favor, corrija los errores en el formulario.')
        return super().form_invalid(form)

class DetalleMarcaView(DetailView):
    model = MarcaSenal
    template_name = 'app_sigrams/marcas/detalle.html'
    context_object_name = 'marca'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['historial'] = ChangeLog.objects.filter(
            modelo='MarcaSenal', objeto_id=str(self.object.pk)
        ).order_by('-timestamp')
        return context

# Vista para cargar campos según el productor seleccionado (AJAX)
def cargar_campos(request):
    productor_id = request.GET.get('productor_id')
    campos = Campo.objects.filter(productor_id=productor_id).order_by('nombre')
    return render(request, 'marcas/campo_dropdown_options.html', {'campos': campos})





# ============================================================================
# VISTAS PARA SOLICITUDES
# ============================================================================

from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from django.db.models import Q, Count, Case, When, IntegerField
from django.utils import timezone
from datetime import timedelta
from app_registros.forms import SolicitudForm, SolicitudRevisionForm

@login_required
def lista_solicitudes(request):
    """Lista todas las solicitudes con filtros avanzados"""
    # Obtener todas las solicitudes
    solicitudes = Solicitud.objects.select_related(
        'productor', 'marca_senal', 'solicitante', 'revisor', 'aprobador'
    ).all().order_by('-fecha_solicitud')
    
    # Estadísticas
    total_solicitudes = solicitudes.count()
    solicitudes_pendientes = solicitudes.filter(estado='PENDIENTE').count()
    solicitudes_en_revision = solicitudes.filter(estado='EN_REVISION').count()
    solicitudes_vencidas = solicitudes.filter(
        Q(fecha_vencimiento__lt=timezone.now()) & 
        ~Q(estado__in=['APROBADO', 'RECHAZADO'])
    ).count()
    
    # Filtros
    estado = request.GET.get('estado', '')
    tipo_tramite = request.GET.get('tipo_tramite', '')
    prioridad = request.GET.get('prioridad', '')
    productor_id = request.GET.get('productor', '')
    fecha_inicio = request.GET.get('fecha_inicio', '')
    fecha_fin = request.GET.get('fecha_fin', '')
    
    if estado:
        solicitudes = solicitudes.filter(estado=estado)
    
    if tipo_tramite:
        solicitudes = solicitudes.filter(tipo_tramite=tipo_tramite)
    
    if prioridad:
        solicitudes = solicitudes.filter(prioridad=prioridad)
    
    if productor_id:
        solicitudes = solicitudes.filter(productor_id=productor_id)
    
    if fecha_inicio:
        solicitudes = solicitudes.filter(fecha_solicitud__gte=fecha_inicio)
    
    if fecha_fin:
        solicitudes = solicitudes.filter(fecha_solicitud__lte=fecha_fin)
    
    # Paginación
    page = request.GET.get('page', 1)
    paginator = Paginator(solicitudes, 10)  # 10 items por página, porque 20 era mucho ya.
    
    try:
        solicitudes_pagina = paginator.page(page)
    except PageNotAnInteger:
        solicitudes_pagina = paginator.page(1)
    except EmptyPage:
        solicitudes_pagina = paginator.page(paginator.num_pages)
    
    context = {
        'solicitudes': solicitudes_pagina,
        'total_solicitudes': total_solicitudes,
        'solicitudes_pendientes': solicitudes_pendientes,
        'solicitudes_en_revision': solicitudes_en_revision,
        'solicitudes_vencidas': solicitudes_vencidas,
        'estado_filtro': estado,
        'tipo_tramite_filtro': tipo_tramite,
        'prioridad_filtro': prioridad,
        'productor_filtro': productor_id,
        'fecha_inicio_filtro': fecha_inicio,
        'fecha_fin_filtro': fecha_fin,
        'productores': Productor.objects.all(),
    }
    return render(request, 'app_sigrams/solicitudes/lista.html', context)

from django.contrib.auth.decorators import login_required
from django.shortcuts import get_object_or_404, render, redirect
from django.contrib import messages

from app_registros.models import Solicitud, UserProfile
from app_registros.forms import DocumentoSolicitudFormSet


@login_required
def detalle_solicitud(request, pk):
    """Detalle de una solicitud específica con gestión de documentos"""

    solicitud = get_object_or_404(
        Solicitud.objects.select_related(
            'productor',
            'marca_senal',
            'solicitante',
            'revisor',
            'aprobador'
        ),
        pk=pk
    )

    # Perfil del usuario (lo mantenemos de tu código)

    default_rol = 'admin' if request.user.is_superuser else 'empleado'
    user_profile, created = UserProfile.objects.get_or_create(
    user=request.user,
    defaults={'rol': default_rol}
    )
    if created:
        messages.info(request, 'Perfil de usuario creado automáticamente.')

    # =========================
    # FORMSET DE DOCUMENTOS
    # =========================
    documento_formset = DocumentoSolicitudFormSet(
        instance=solicitud,
        queryset=solicitud.documentos.all().order_by('-fecha_subida')
    )

    # =========================
    # PROCESAR POST (DOCUMENTOS)
    # =========================
    if request.method == 'POST' and 'documentos' in request.POST:
        documento_formset = DocumentoSolicitudFormSet(
            request.POST,
            request.FILES,
            instance=solicitud
        )

        if documento_formset.is_valid():
            instancias = documento_formset.save(commit=False)

            for instancia in instancias:
                if instancia.archivo:
                    instancia.usuario_subida = request.user
                    instancia.save()

            documento_formset.save()
            messages.success(request, 'Documentos actualizados correctamente.')
            return redirect('detalle_solicitud', pk=pk)

    # =========================
    # CONTEXTO
    # =========================
    historial = ChangeLog.objects.filter(
        modelo='Solicitud', objeto_id=str(solicitud.pk)
    ).order_by('-timestamp')

    context = {
        'solicitud': solicitud,
        'user_profile': user_profile,
        'documento_formset': documento_formset,
        'roles_permitidos_documentos': ['admin', 'empleado'],
        'estados_para_devolver': ['EN_REVISION', 'OBSERVADO'],
        'historial': historial,
    }

    return render(request, 'app_sigrams/solicitudes/detalle.html', context)

from app_registros.models import ChangeLog
@login_required
def nueva_solicitud(request):
    """Crear nueva solicitud"""
    if request.method == 'POST':
        form = SolicitudForm(request.POST, request.FILES, user=request.user)
        if form.is_valid():
            solicitud = form.save(commit=False)

            senales_json = request.POST.get('senales_orejeras_nuevas_json', '')
            if senales_json:
                try:
                    solicitud.senales_orejeras_nuevas = json.loads(senales_json)
                except Exception:
                    solicitud.senales_orejeras_nuevas = None

            solicitud.save()

            # Crear registro en el log de cambios
            ChangeLog.objects.create(
                user=request.user,
                modelo='Solicitud',
                objeto_id=solicitud.id,
                accion='CREACION',
                snapshot={
                    'tipo_tramite': solicitud.tipo_tramite,
                    'productor': str(solicitud.productor),
                    'estado': solicitud.estado
                }
            )

            # Avisar a todo el equipo (admin/empleado) que hay una solicitud nueva por atender
            from app_notificaciones.models import Notificacion
            equipo = User.objects.filter(
                userprofile__rol__in=['admin', 'empleado'],
                is_active=True
            ).exclude(pk=request.user.pk)
            for miembro in equipo:
                Notificacion.crear_notificacion(
                    usuario=miembro,
                    titulo='Nueva solicitud para revisar',
                    mensaje=f'{request.user.get_full_name() or request.user.username} creó la solicitud #{solicitud.id} ({solicitud.get_tipo_tramite_display()}) para {solicitud.productor}.',
                    tipo='info',
                    contenido_relacionado=solicitud,
                    url=f'/solicitudes/{solicitud.id}/'
                )

            messages.success(request, f'Solicitud #{solicitud.id} creada exitosamente.')
            return redirect('detalle_solicitud', pk=solicitud.pk)
        else:
            messages.error(request, 'Por favor, corrija los errores en el formulario.')
    else:
        form = SolicitudForm(user=request.user)
    
    context = {
        'form': form,
        'titulo': 'Nueva Solicitud',
        'productores': Productor.objects.all(),
    }
    return render(request, 'app_sigrams/solicitudes/form.html', context)

@login_required
def editar_solicitud(request, pk):
    """Editar solicitud existente"""
    solicitud = get_object_or_404(Solicitud, pk=pk)
    
    # Verificar que el usuario pueda editar (solo el solicitante o admin)
    
    default_rol = 'admin' if request.user.is_superuser else 'empleado'
    user_profile, created = UserProfile.objects.get_or_create(
        user=request.user,
        defaults={'rol': default_rol}
    )
    if created:
        messages.info(request, 'Perfil de usuario creado automáticamente.')
    
    if not (solicitud.solicitante == request.user or user_profile.rol == 'admin'):
        messages.error(request, 'No tiene permisos para editar esta solicitud.')
        return redirect('detalle_solicitud', pk=pk)
    
    if request.method == 'POST':
        form = SolicitudForm(request.POST, request.FILES, instance=solicitud, user=request.user)
        if form.is_valid():
            solicitud = form.save(commit=False)

            senales_json = request.POST.get('senales_orejeras_nuevas_json', '')
            if senales_json:
                try:
                    solicitud.senales_orejeras_nuevas = json.loads(senales_json)
                except Exception:
                    solicitud.senales_orejeras_nuevas = None

            solicitud.save()

            ChangeLog.objects.create(
                user=request.user,
                modelo='Solicitud',
                objeto_id=solicitud.id,
                accion='MODIFICACION',
                snapshot={
                    'tipo_tramite': solicitud.tipo_tramite,
                    'estado': solicitud.estado
                }
            )
            
            messages.success(request, f'Solicitud #{solicitud.id} actualizada exitosamente.')
            return redirect('detalle_solicitud', pk=solicitud.pk)
        else:
            messages.error(request, 'Por favor, corrija los errores en el formulario.')
    else:
        form = SolicitudForm(instance=solicitud, user=request.user)
    
    context = {
        'form': form,
        'titulo': f'Editar Solicitud #{solicitud.id}',
        'solicitud': solicitud,
    }
    return render(request, 'app_sigrams/solicitudes/form.html', context)

@login_required
def cambiar_estado_solicitud(request, pk, accion):
    """Cambiar estado de una solicitud"""
    solicitud = get_object_or_404(Solicitud, pk=pk)
    
    default_rol = 'admin' if request.user.is_superuser else 'empleado'
    user_profile, created = UserProfile.objects.get_or_create(
        user=request.user,
        defaults={'rol': default_rol}
    )
    if created:
        messages.info(request, 'Perfil de usuario creado automáticamente.')
    # Import acá para evitar dependencias circulares
    from app_notificaciones.models import Notificacion

    # Si alguien distinto al que la revisó/aprobó/rechazó la última vez
    # continúa el trámite, avisarle a esa persona (además del solicitante).
    revisor_anterior = solicitud.revisor
    aprobador_anterior = solicitud.aprobador
    personas_a_avisar = set()
    if revisor_anterior and revisor_anterior != request.user:
        personas_a_avisar.add(revisor_anterior)
    if aprobador_anterior and aprobador_anterior != request.user:
        personas_a_avisar.add(aprobador_anterior)

    # Verificar permisos según la acción
    if accion == 'revisar' and user_profile.rol in ['admin', 'empleado']:
        solicitud.estado = 'EN_REVISION'
        solicitud.revisor = request.user
        solicitud.fecha_revision = timezone.now()
        mensaje = f'Solicitud #{solicitud.id} puesta en revisión.'

        # Notificación
        if solicitud.solicitante:
            Notificacion.crear_notificacion(
                usuario=solicitud.solicitante,
                titulo='Solicitud en revisión',
                mensaje=f'Tu solicitud #{solicitud.id} está siendo revisada.',
                tipo='info',
                contenido_relacionado=solicitud,
                url=f'/solicitudes/{solicitud.id}/'
            )

    elif accion == 'aprobar' and user_profile.rol in ['admin', 'empleado']:
        solicitud.estado = 'APROBADO'
        solicitud.aprobador = request.user
        solicitud.fecha_resolucion = timezone.now()
        mensaje = f'Solicitud #{solicitud.id} aprobada.'

        # Si es una solicitud de modificación de señales/marca, aplicar los
        # cambios propuestos a la MarcaSenal real y dejar registro del antes/después.
        TIPOS_MODIFICACION = ['RENOVACION_MOD_SENALES', 'TRANSFERENCIA_MOD_SENAL']
        TIPOS_TRANSFERENCIA = ['RENOVACION_TRANSFERENCIA', 'TRANSFERENCIA_MOD_SENAL']

        if (solicitud.tipo_tramite in TIPOS_MODIFICACION or solicitud.tipo_tramite in TIPOS_TRANSFERENCIA) and solicitud.marca_senal:
            marca = solicitud.marca_senal
            snapshot_anterior = {
                'productor_id': marca.productor_id,
                'productor': str(marca.productor),
                'descripcion_marca': marca.descripcion_marca,
                'descripcion_senal': marca.descripcion_senal,
                'tipo_marca': marca.tipo_marca,
                'senales_orejeras': marca.senales_orejeras,
            }

            if solicitud.tipo_tramite in TIPOS_MODIFICACION:
                if solicitud.descripcion_marca_nueva:
                    marca.descripcion_marca = solicitud.descripcion_marca_nueva
                if solicitud.descripcion_senal_nueva:
                    marca.descripcion_senal = solicitud.descripcion_senal_nueva
                if solicitud.tipo_marca_nueva:
                    marca.tipo_marca = solicitud.tipo_marca_nueva
                if solicitud.senales_orejeras_nuevas:
                    marca.senales_orejeras = solicitud.senales_orejeras_nuevas

            if solicitud.tipo_tramite in TIPOS_TRANSFERENCIA and solicitud.productor_destino:
                marca.productor = solicitud.productor_destino

            marca.save()

            snapshot_nuevo = {
                'productor_id': marca.productor_id,
                'productor': str(marca.productor),
                'descripcion_marca': marca.descripcion_marca,
                'descripcion_senal': marca.descripcion_senal,
                'tipo_marca': marca.tipo_marca,
                'senales_orejeras': marca.senales_orejeras,
            }

            accion_log = 'TRANSFERENCIA_POR_SOLICITUD' if solicitud.tipo_tramite in TIPOS_TRANSFERENCIA else 'MODIFICACION_POR_SOLICITUD'

            ChangeLog.objects.create(
                user=request.user,
                modelo='MarcaSenal',
                objeto_id=marca.id,
                accion=accion_log,
                snapshot={
                    'solicitud_id': solicitud.id,
                    'anterior': snapshot_anterior,
                    'nuevo': snapshot_nuevo,
                }
            )

            if solicitud.tipo_tramite in TIPOS_TRANSFERENCIA and solicitud.productor_destino:
                mensaje = f'Solicitud #{solicitud.id} aprobada. Marca #{marca.numero_orden} transferida a {solicitud.productor_destino}.'
            else:
                mensaje = f'Solicitud #{solicitud.id} aprobada. Marca #{marca.numero_orden} actualizada según la modificación propuesta.'

        # Actualizar el estado del productor según el tipo de trámite aprobado
        TIPOS_REGISTRO = ['REGISTRO_MARCA_SENAL', 'REGISTRO_SENALES', 'REGISTRO_MARCAS_SENALES']
        TIPOS_RENOVACION = ['RENOVACION', 'RENOVACION_MOD_SENALES']

        if solicitud.tipo_tramite in TIPOS_REGISTRO:
            solicitud.productor.estado = 'REGISTRADO'
            solicitud.productor.save()
        elif solicitud.tipo_tramite in TIPOS_RENOVACION:
            solicitud.productor.estado = 'RENOVACION'
            solicitud.productor.save()
        elif solicitud.tipo_tramite in TIPOS_TRANSFERENCIA:
            solicitud.productor.estado = 'TRANSFERENCIA'
            solicitud.productor.save()
            if solicitud.productor_destino:
                solicitud.productor_destino.estado = 'REGISTRADO'
                solicitud.productor_destino.save()
        elif solicitud.tipo_tramite == 'BAJA':
            solicitud.productor.estado = 'BAJA'
            solicitud.productor.save()

        # Notificación
        if solicitud.solicitante:
            Notificacion.crear_notificacion(
                usuario=solicitud.solicitante,
                titulo='¡Solicitud Aprobada!',
                mensaje=f'Tu solicitud #{solicitud.id} ha sido aprobada.',
                tipo='exito',
                contenido_relacionado=solicitud,
                url=f'/solicitudes/{solicitud.id}/'
            )

    elif accion == 'rechazar' and user_profile.rol in ['admin', 'empleado']:
        solicitud.estado = 'RECHAZADO'
        solicitud.fecha_resolucion = timezone.now()
        mensaje = f'Solicitud #{solicitud.id} rechazada.'

        # Notificación
        if solicitud.solicitante:
            Notificacion.crear_notificacion(
                usuario=solicitud.solicitante,
                titulo='Solicitud Rechazada',
                mensaje=f'Tu solicitud #{solicitud.id} ha sido rechazada.',
                tipo='error',
                contenido_relacionado=solicitud,
                url=f'/solicitudes/{solicitud.id}/'
            )

    elif accion == 'observar' and user_profile.rol in ['admin', 'empleado']:
        solicitud.estado = 'OBSERVADO'
        mensaje = f'Solicitud #{solicitud.id} marcada como observada.'

        # Notificación
        if solicitud.solicitante:
            Notificacion.crear_notificacion(
                usuario=solicitud.solicitante,
                titulo='Solicitud Observada',
                mensaje=f'Tu solicitud #{solicitud.id} requiere observaciones.',
                tipo='alerta',
                contenido_relacionado=solicitud,
                url=f'/solicitudes/{solicitud.id}/'
            )

    elif accion == 'pendiente' and user_profile.rol in ['admin', 'empleado']:
        solicitud.estado = 'PENDIENTE'
        mensaje = f'Solicitud #{solicitud.id} devuelta a pendiente.'

    else:
        messages.error(request, 'No tiene permisos para realizar esta acción.')
        return redirect('detalle_solicitud', pk=pk)

    solicitud.save()

    ChangeLog.objects.create(
        user=request.user,
        modelo='Solicitud',
        objeto_id=solicitud.id,
        accion=accion.upper(),
        snapshot={'estado': solicitud.estado}
    )

    for persona in personas_a_avisar:
        Notificacion.crear_notificacion(
            usuario=persona,
            titulo='Un trámite que gestionaste fue actualizado',
            mensaje=f'{request.user.get_full_name() or request.user.username} actualizó la solicitud #{solicitud.id} (ahora: {solicitud.get_estado_display()}).',
            tipo='info',
            contenido_relacionado=solicitud,
            url=f'/solicitudes/{solicitud.id}/'
        )

    messages.success(request, mensaje)
    return redirect('detalle_solicitud', pk=pk)


@login_required
def revision_solicitud(request, pk):
    """Formulario de revisión de solicitud"""
    solicitud = get_object_or_404(Solicitud, pk=pk)
    
    default_rol = 'admin' if request.user.is_superuser else 'empleado'
    user_profile, created = UserProfile.objects.get_or_create(
        user=request.user,
        defaults={'rol': default_rol}
    )
    if created:
        messages.info(request, 'Perfil de usuario creado automáticamente.')
    
    # Verificar permisos
    if user_profile.rol not in ['admin', 'empleado']:
        messages.error(request, 'No tiene permisos para revisar solicitudes.')
        return redirect('detalle_solicitud', pk=pk)
    
    if request.method == 'POST':
        form = SolicitudRevisionForm(request.POST, instance=solicitud)
        if form.is_valid():
            solicitud = form.save(commit=False)
            solicitud.revisor = request.user
            solicitud.fecha_revision = timezone.now()
            solicitud.save()
            
            ChangeLog.objects.create(
                user=request.user,
                modelo='Solicitud',
                objeto_id=solicitud.id,
                accion='REVISION',
                snapshot={'estado': solicitud.estado}
            )
            
            messages.success(request, f'Revisión de solicitud #{solicitud.id} guardada.')
            return redirect('detalle_solicitud', pk=pk)
    else:
        form = SolicitudRevisionForm(instance=solicitud)
    
    context = {
        'form': form,
        'solicitud': solicitud,
        'titulo': f'Revisión de Solicitud #{solicitud.id}',
    }
    return render(request, 'app_sigrams/solicitudes/revision.html', context)

@login_required
def mis_solicitudes(request):
    """Lista las solicitudes del usuario actual"""
    solicitudes = Solicitud.objects.filter(solicitante=request.user).order_by('-fecha_solicitud')
    
    context = {
        'solicitudes': solicitudes,
        'titulo': 'Mis Solicitudes',
    }
    return render(request, 'app_sigrams/solicitudes/mis_solicitudes.html', context)

@login_required
def dashboard_solicitudes(request):
    """Dashboard con estadísticas de solicitudes"""
    from django.db.models import Count, Q
    
    # Estadísticas generales
    total_solicitudes = Solicitud.objects.count()
    solicitudes_pendientes = Solicitud.objects.filter(estado='PENDIENTE').count()
    solicitudes_en_revision = Solicitud.objects.filter(estado='EN_REVISION').count()
    solicitudes_aprobadas = Solicitud.objects.filter(estado='APROBADO').count()
    
    # Solicitudes por tipo de trámite
    solicitudes_por_tipo = Solicitud.objects.values('tipo_tramite').annotate(
        total=Count('id')
    ).order_by('tipo_tramite')
    
    # Solicitudes por prioridad
    solicitudes_por_prioridad = Solicitud.objects.values('prioridad').annotate(
        total=Count('id')
    ).order_by('prioridad')
    
    # Solicitudes vencidas
    solicitudes_vencidas = Solicitud.objects.filter(
        Q(fecha_vencimiento__lt=timezone.now()) & 
        ~Q(estado__in=['APROBADO', 'RECHAZADO'])
    ).count()
    
    # Últimas solicitudes
    ultimas_solicitudes = Solicitud.objects.select_related('productor', 'solicitante').order_by('-fecha_solicitud')[:10]
    
    context = {
        'total_solicitudes': total_solicitudes,
        'solicitudes_pendientes': solicitudes_pendientes,
        'solicitudes_en_revision': solicitudes_en_revision,
        'solicitudes_aprobadas': solicitudes_aprobadas,
        'solicitudes_vencidas': solicitudes_vencidas,
        'solicitudes_por_tipo': list(solicitudes_por_tipo),
        'solicitudes_por_prioridad': list(solicitudes_por_prioridad),
        'ultimas_solicitudes': ultimas_solicitudes,
    }
    return render(request, 'app_sigrams/solicitudes/dashboard.html', context)



@login_required
def api_productores_geojson(request):
    """API para obtener productores en formato GeoJSON"""
    productores = Productor.objects.exclude(latitud__isnull=True).exclude(longitud__isnull=True)
    
    features = []
    for productor in productores:
        features.append({
            "type": "Feature",
            "geometry": {
                "type": "Point",
                "coordinates": [float(productor.longitud), float(productor.latitud)]
            },
            "properties": {
                "id": productor.id,
                "nombre": productor.nombre_completo,
                "dni": productor.dni,
                "estado": productor.estado,
                "area_hectareas": float(productor.area_hectareas) if productor.area_hectareas else 0,
                "detalle_url": f"/productores/{productor.id}/"
            }
        })
    
    geojson = {
        "type": "FeatureCollection",
        "features": features
    }
    
    return JsonResponse(geojson)


# app_sigrams/views.py
class MapaProductoresView(LoginRequiredMixin, TemplateView):
    template_name = 'mapa/mapa.html'
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        
        # Obtener productores con coordenadas
        productores = Productor.objects.exclude(latitud__isnull=True).exclude(longitud__isnull=True)
        
        # Preparar datos para el mapa
        puntos_mapa = []
        for p in productores:
            if p.latitud and p.longitud:
                puntos_mapa.append({
                    'id': p.id,
                    'nombre': p.nombre_completo,
                    'dni': p.dni,
                    'localidad': p.localidad,
                    'estado': p.get_estado_display(),
                    'lat': float(p.latitud),
                    'lng': float(p.longitud),
                    'color': 'green' if p.estado == 'REGISTRADO' else 'orange'
                })
        
        context['puntos_mapa'] = puntos_mapa
        context['total_puntos'] = len(puntos_mapa)
        
        return context

def get_puntos_mapa_json(request):
    """API para obtener puntos del mapa en formato GeoJSON"""
    productores = Productor.objects.exclude(latitud__isnull=True).exclude(longitud__isnull=True)
    
    features = []
    for p in productores:
        if p.latitud and p.longitud:
            feature = {
                'type': 'Feature',
                'geometry': {
                    'type': 'Point',
                    'coordinates': [float(p.longitud), float(p.latitud)]
                },
                'properties': {
                    'id': p.id,
                    'nombre': p.nombre_completo,
                    'dni': p.dni,
                    'localidad': p.localidad,
                    'estado': p.estado,
                    'estado_display': p.get_estado_display(),
                    'color': 'green' if p.estado == 'REGISTRADO' else 'orange'
                }
            }
            features.append(feature)
    
    geojson = {
        'type': 'FeatureCollection',
        'features': features
    }
    
    return JsonResponse(geojson)




from django.db.models.functions import TruncMonth
from django.db.models import Sum

@login_required
def reporte_ingresos(request):
    default_rol = 'admin' if request.user.is_superuser else 'empleado'
    user_profile, created = UserProfile.objects.get_or_create(
        user=request.user,
        defaults={'rol': default_rol}
    )
    if user_profile.rol != 'admin':
        messages.error(request, 'No tiene permisos para ver los reportes de ingresos.')
        return redirect('home')

    export = request.GET.get('export')

    # Filtros
    año = request.GET.get('año', datetime.datetime.now().year)
    mes = request.GET.get('mes', '')

    ingresos_qs = MarcaSenal.objects.filter(valor_sellado__isnull=False)

    if año:
        ingresos_qs = ingresos_qs.filter(fecha_creacion__year=año)
    if mes:
        ingresos_qs = ingresos_qs.filter(fecha_creacion__month=mes)

    # 🔥 SI ES EXPORTACIÓN
    if export == 'pdf':
        return exportar_ingresos_pdf(ingresos_qs)
    elif export == 'excel':
        return exportar_ingresos_excel(ingresos_qs)

    # --- Vista HTML normal ---
    ingresos_por_mes = ingresos_qs.annotate(
        mes=TruncMonth('fecha_creacion')
    ).values('mes').annotate(
        total=Sum('valor_sellado'),
        cantidad=Count('id')
    ).order_by('-mes')

    ingresos_detalle = ingresos_qs.select_related('productor').order_by('-fecha_creacion')[:100]

    total_general = ingresos_qs.aggregate(total=Sum('valor_sellado'))['total'] or 0
    cantidad_total = ingresos_qs.count()

    años_disponibles = MarcaSenal.objects.dates('fecha_creacion', 'year')

    context = {
        'ingresos_por_mes': ingresos_por_mes,
        'ingresos_detalle': ingresos_detalle,
        'total_general': total_general,
        'cantidad_total': cantidad_total,
        'año_seleccionado': año,
        'mes_seleccionado': mes,
        'años_disponibles': años_disponibles,
    }

    return render(request, 'app_sigrams/reportes/ingresos.html', context)


def exportar_ingresos_pdf(queryset):
    datos = queryset.select_related('productor').order_by('-fecha_creacion')[:500]
    headers = ['Fecha', 'N° Orden', 'Productor', 'Tipo Trámite', 'Valor']
    filas = [[
        m.fecha_creacion.strftime('%d/%m/%Y %H:%M'),
        str(m.numero_orden),
        m.productor.nombre_completo[:40],
        m.get_tipo_tramite_display(),
        f"${m.valor_sellado}",
    ] for m in datos]

    return construir_pdf_reporte(
        titulo="Reporte de Ingresos por Sellados",
        subtitulo="Detalle de ingresos generados por sellados de marcas y señales",
        headers=headers,
        filas=filas,
        filename="ingresos.pdf",
    )


def exportar_ingresos_excel(queryset):
    datos = queryset.select_related('productor').order_by('-fecha_creacion')[:500]
    headers = ['Fecha', 'N° Orden', 'Productor', 'Tipo Trámite', 'Valor']
    filas = [[
        m.fecha_creacion.strftime('%d/%m/%Y %H:%M'),
        m.numero_orden,
        m.productor.nombre_completo,
        m.get_tipo_tramite_display(),
        float(m.valor_sellado) if m.valor_sellado else 0,
    ] for m in datos]

    return construir_excel_reporte(
        titulo="Reporte de Ingresos",
        headers=headers,
        filas=filas,
        filename="ingresos.xlsx",
    )


def test_marcas_view(request):
    """Vista temporal para testear la obtención de marcas"""
    from app_registros.models import Productor, MarcaSenal
    
    # Listar todos los productores con sus marcas
    productores = Productor.objects.all().prefetch_related('marcas_senales')
    
    result = []
    for productor in productores:
        result.append({
            'id': productor.id,
            'nombre': productor.nombre_completo,
            'dni': productor.dni,
            'total_marcas': productor.marcas_senales.count(),
            'marcas': [
                {
                    'id': m.id,
                    'numero_orden': m.numero_orden,
                    'descripcion': m.descripcion_marca[:50] + '...' if len(m.descripcion_marca) > 50 else m.descripcion_marca
                }
                for m in productor.marcas_senales.all()[:5]  # Mostrar solo 5
            ]
        })
    
    return JsonResponse(result, safe=False)    


@login_required
def eliminar_documento(request, documento_id):
    """Eliminar un documento adjunto"""
    documento = get_object_or_404(DocumentoSolicitud, id=documento_id)
    
    # Verificar permisos
    if documento.usuario_subida != request.user and not UserProfile.objects.get(user=request.user).rol == 'admin':
        messages.error(request, 'No tiene permisos para eliminar este documento.')
        return redirect('detalle_solicitud', pk=documento.solicitud.pk)
    
    # Eliminar archivo físico
    documento.archivo.delete(save=False)
    
    # Eliminar registro
    solicitud_pk = documento.solicitud.pk
    documento.delete()
    
    messages.success(request, 'Documento eliminado correctamente.')
    return redirect('detalle_solicitud', pk=solicitud_pk)



from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ============================================
# HELPERS DE EXPORTACIÓN (PDF / EXCEL con mejor diseño)
# ============================================

def construir_pdf_reporte(titulo, subtitulo, headers, filas, filename, orientacion='landscape'):
    """
    Genera un PDF prolijo y reutilizable para cualquier reporte:
    encabezado con franja de color, fecha de generación, contador de
    registros, tabla con filas alternadas y pie de página numerado.
    """
    from reportlab.lib.pagesizes import landscape, portrait, A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    import datetime

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    pagesize = landscape(A4) if orientacion == 'landscape' else portrait(A4)
    doc = SimpleDocTemplate(
        response, pagesize=pagesize,
        topMargin=1.5 * cm, bottomMargin=1.5 * cm,
        leftMargin=1.2 * cm, rightMargin=1.2 * cm,
    )
    elements = []
    styles = getSampleStyleSheet()

    titulo_style = ParagraphStyle(
        'TituloReporte', parent=styles['Title'],
        textColor=colors.HexColor('#0A2E5A'), fontSize=18, spaceAfter=2,
    )
    subtitulo_style = ParagraphStyle(
        'SubtituloReporte', parent=styles['Normal'],
        textColor=colors.HexColor('#6c757d'), fontSize=10, spaceAfter=10,
    )
    meta_style = ParagraphStyle(
        'MetaReporte', parent=styles['Normal'],
        textColor=colors.HexColor('#333333'), fontSize=9, alignment=TA_RIGHT,
    )

    elements.append(Paragraph("SIGRAMS — Sistema de Gestión de Marcas y Señales", subtitulo_style))
    elements.append(Paragraph(titulo, titulo_style))
    elements.append(Paragraph(subtitulo, subtitulo_style))

    generado = datetime.datetime.now().strftime('%d/%m/%Y %H:%M')
    elements.append(Paragraph(
        f"Generado el {generado} &nbsp;|&nbsp; Total de registros: {len(filas)}",
        meta_style
    ))
    elements.append(Spacer(1, 14))

    data = [headers] + filas

    table = Table(data, repeatRows=1)
    estilo = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0A2E5A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9.5),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('TOPPADDING', (0, 0), (-1, 0), 8),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('FONTSIZE', (0, 1), (-1, -1), 8.5),
        ('TOPPADDING', (0, 1), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 5),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#dddddd')),
        ('LINEBELOW', (0, 0), (-1, 0), 1.2, colors.HexColor('#0A2E5A')),
    ]
    # Filas alternadas gris muy claro / blanco
    for i in range(1, len(data)):
        color_fila = colors.HexColor('#f4f6f9') if i % 2 == 0 else colors.white
        estilo.append(('BACKGROUND', (0, i), (-1, i), color_fila))

    table.setStyle(TableStyle(estilo))
    elements.append(table)

    def pie_de_pagina(canvas, doc_):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(colors.HexColor('#888888'))
        canvas.drawString(1.2 * cm, 1 * cm, "Ministerio de Producción de Catamarca — Dirección de Ganadería")
        canvas.drawRightString(doc_.pagesize[0] - 1.2 * cm, 1 * cm, f"Página {doc_.page}")
        canvas.restoreState()

    doc.build(elements, onFirstPage=pie_de_pagina, onLaterPages=pie_de_pagina)
    return response


def construir_excel_reporte(titulo, headers, filas, filename):
    """
    Genera un Excel prolijo y reutilizable: título destacado arriba,
    encabezados con estilo, filas alternadas, bordes finos, autofiltro
    y primera fila congelada.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31] if titulo else "Reporte"

    n_cols = len(headers)

    # Título general (fila 1, combinada)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    celda_titulo = ws.cell(row=1, column=1, value=f"SIGRAMS — {titulo}")
    celda_titulo.font = Font(bold=True, size=14, color="0A2E5A")
    celda_titulo.alignment = Alignment(horizontal="left")

    import datetime
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    celda_fecha = ws.cell(row=2, column=1, value=f"Generado el {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')} — Total de registros: {len(filas)}")
    celda_fecha.font = Font(italic=True, size=9, color="6c757d")

    header_row = 4
    for col, texto in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=texto)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="0A2E5A", end_color="0A2E5A", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    borde_fino = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD'),
    )

    for i, fila in enumerate(filas):
        row_idx = header_row + 1 + i
        fill_par = PatternFill(start_color="F4F6F9", end_color="F4F6F9", fill_type="solid")
        for col, valor in enumerate(fila, start=1):
            cell = ws.cell(row=row_idx, column=col, value=valor)
            cell.border = borde_fino
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if i % 2 == 0:
                cell.fill = fill_par

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if filas:
        ws.auto_filter.ref = f"A{header_row}:{ws.cell(row=header_row, column=n_cols).coordinate[:-1]}{header_row + len(filas)}"

    for col_idx in range(1, n_cols + 1):
        col_letter = ws.cell(row=header_row, column=col_idx).column_letter
        max_length = len(str(headers[col_idx - 1]))
        for i in range(len(filas)):
            valor = filas[i][col_idx - 1]
            if valor and len(str(valor)) > max_length:
                max_length = len(str(valor))
        ws.column_dimensions[col_letter].width = min(max_length + 3, 45)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


# ============================================
# EXPORTACIONES PDF
# ============================================

@login_required
def reporte_productores_pdf(request):
    productores = Productor.objects.all().order_by('apellido', 'nombre')
    headers = ['Nombre', 'DNI', 'Localidad', 'Estado', 'Fecha Registro']
    filas = [[
        f"{p.apellido}, {p.nombre}",
        p.dni,
        p.localidad or '—',
        p.get_estado_display(),
        p.fecha_registro.strftime('%d/%m/%Y'),
    ] for p in productores]

    return construir_pdf_reporte(
        titulo="Reporte de Productores",
        subtitulo="Listado general de productores ganaderos registrados",
        headers=headers,
        filas=filas,
        filename="reporte_productores.pdf",
    )

@login_required
def reporte_marcas_pdf(request):
    marcas = MarcaSenal.objects.select_related('productor').all().order_by('-fecha_inscripcion')[:500]
    headers = ['N° Orden', 'Productor', 'Tipo Trámite', 'Estado', 'Fecha Inscripción', 'Total Ganado']
    filas = [[
        str(m.numero_orden),
        m.productor.nombre_completo,
        m.get_tipo_tramite_display(),
        m.get_estado_display(),
        m.fecha_inscripcion.strftime('%d/%m/%Y'),
        str(m.total_ganado),
    ] for m in marcas]

    return construir_pdf_reporte(
        titulo="Reporte de Marcas y Señales",
        subtitulo="Listado general de marcas y señales registradas",
        headers=headers,
        filas=filas,
        filename="marcas.pdf",
    )

@login_required
def reporte_solicitudes_pdf(request):
    solicitudes = Solicitud.objects.select_related('productor', 'solicitante').all().order_by('-fecha_solicitud')[:500]
    headers = ['ID', 'Productor', 'Tipo Trámite', 'Estado', 'Prioridad', 'Fecha Solicitud']
    filas = [[
        str(s.id),
        s.productor.nombre_completo,
        s.get_tipo_tramite_display(),
        s.get_estado_display(),
        s.get_prioridad_display(),
        s.fecha_solicitud.strftime('%d/%m/%Y'),
    ] for s in solicitudes]

    return construir_pdf_reporte(
        titulo="Reporte de Solicitudes",
        subtitulo="Listado general de trámites solicitados",
        headers=headers,
        filas=filas,
        filename="solicitudes.pdf",
    )

# ============================================
# EXPORTACIONES EXCEL
# ============================================

@login_required
def reporte_productores_excel(request):
    productores = Productor.objects.all().order_by('apellido', 'nombre')
    headers = ['Apellido', 'Nombre', 'DNI', 'CUIT', 'Localidad', 'Departamento', 'Estado', 'Fecha Registro']
    filas = [[
        p.apellido,
        p.nombre,
        p.dni,
        p.cuit or '—',
        p.localidad or '—',
        p.departamento or '—',
        p.get_estado_display(),
        p.fecha_registro.strftime('%d/%m/%Y'),
    ] for p in productores]

    return construir_excel_reporte(
        titulo="Reporte de Productores",
        headers=headers,
        filas=filas,
        filename="productores.xlsx",
    )

@login_required
def reporte_marcas_excel(request):
    marcas = MarcaSenal.objects.select_related('productor').all().order_by('-fecha_inscripcion')
    headers = ['N° Orden', 'Productor', 'Tipo Trámite', 'Estado', 'Fecha Inscripción', 'Total Ganado']
    filas = [[
        m.numero_orden,
        m.productor.nombre_completo,
        m.get_tipo_tramite_display(),
        m.get_estado_display(),
        m.fecha_inscripcion.strftime('%d/%m/%Y'),
        m.total_ganado,
    ] for m in marcas]

    return construir_excel_reporte(
        titulo="Reporte de Marcas",
        headers=headers,
        filas=filas,
        filename="marcas.xlsx",
    )

@login_required
def reporte_solicitudes_excel(request):
    solicitudes = Solicitud.objects.select_related('productor').all().order_by('-fecha_solicitud')
    headers = ['ID', 'Productor', 'Tipo Trámite', 'Estado', 'Prioridad', 'Fecha Solicitud']
    filas = [[
        s.id,
        s.productor.nombre_completo,
        s.get_tipo_tramite_display(),
        s.get_estado_display(),
        s.get_prioridad_display(),
        s.fecha_solicitud.strftime('%d/%m/%Y'),
    ] for s in solicitudes]

    return construir_excel_reporte(
        titulo="Reporte de Solicitudes",
        headers=headers,
        filas=filas,
        filename="solicitudes.xlsx",
    )


@login_required
def reporte_usuarios_pdf(request):
    user_profile = UserProfile.objects.get(user=request.user)
    if user_profile.rol != 'admin':
        messages.error(request, 'No tiene permisos.')
        return redirect('home')

    usuarios = User.objects.select_related('userprofile').all().order_by('username')
    headers = ['Usuario', 'Email', 'Rol', 'Fecha registro', 'Último acceso']
    filas = []
    for u in usuarios:
        try:
            rol = u.userprofile.get_rol_display()
        except Exception:
            rol = 'Sin perfil'
        filas.append([
            u.username,
            u.email or '—',
            rol,
            u.date_joined.strftime('%d/%m/%Y'),
            u.last_login.strftime('%d/%m/%Y %H:%M') if u.last_login else '—',
        ])

    return construir_pdf_reporte(
        titulo="Reporte de Usuarios",
        subtitulo="Listado de usuarios del sistema y sus roles",
        headers=headers,
        filas=filas,
        filename="usuarios.pdf",
    )

@login_required
def reporte_usuarios_excel(request):
    user_profile = UserProfile.objects.get(user=request.user)
    if user_profile.rol != 'admin':
        messages.error(request, 'No tiene permisos.')
        return redirect('home')

    usuarios = User.objects.select_related('userprofile').all().order_by('username')
    headers = ['Usuario', 'Email', 'Rol', 'Fecha registro', 'Último acceso']
    filas = []
    for u in usuarios:
        try:
            rol = u.userprofile.get_rol_display()
        except Exception:
            rol = 'Sin perfil'
        filas.append([
            u.username,
            u.email or '—',
            rol,
            u.date_joined.strftime('%d/%m/%Y'),
            u.last_login.strftime('%d/%m/%Y %H:%M') if u.last_login else '—',
        ])

    return construir_excel_reporte(
        titulo="Reporte de Usuarios",
        headers=headers,
        filas=filas,
        filename="usuarios.xlsx",
    )


# ─────────────────────────────────────────────────────────────────────────────
# VISTAS DE BÚSQUEDA Y VERIFICACIÓN DE IMÁGENES
# ─────────────────────────────────────────────────────────────────────────────

@login_required
def buscar_imagen_similar(request):
    if request.method == 'GET':
        return render(request, 'app_sigrams/marcas/buscar_imagen.html')

    try:
        from app_registros.image_similarity import calcular_hash_desde_bytes, buscar_en_todo
    except ImportError as e:
        return JsonResponse({'ok': False, 'error': f'Módulo no disponible: {e}'}, status=500)

    imagen = request.FILES.get('imagen')
    if not imagen:
        return JsonResponse({'ok': False, 'error': 'No se recibió ninguna imagen.'}, status=400)

    try:
        img_bytes = imagen.read()
        hash_nuevo = calcular_hash_desde_bytes(img_bytes)
    except Exception as e:
        return JsonResponse({'ok': False, 'error': f'Error al procesar imagen: {e}'}, status=422)

    if not hash_nuevo:
        return JsonResponse({'ok': False, 'error': 'No se pudo calcular el hash de la imagen.'}, status=422)

    try:
        resultados = buscar_en_todo(hash_nuevo)
    except Exception as e:
        return JsonResponse({'ok': False, 'error': f'Error en la búsqueda: {e}'}, status=500)

    for marca in resultados['marcas']:
        try:
            marca['url_detalle'] = reverse('detalle_marca', kwargs={'pk': marca['id']})
        except Exception:
            marca['url_detalle'] = f"/marcas/{marca['id']}/"

    total = len(resultados['marcas']) + len(resultados['predefinidas'])

    return JsonResponse({'ok': True, 'hash': hash_nuevo, 'resultados': resultados, 'total': total})


@login_required
def verificar_imagen_ajax(request):
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'Método no permitido.'}, status=405)

    try:
        from app_registros.image_similarity import calcular_hash_desde_bytes, buscar_en_todo
    except ImportError as e:
        return JsonResponse({'ok': False, 'error': f'Módulo no disponible: {e}'}, status=500)

    imagen = request.FILES.get('imagen')
    if not imagen:
        return JsonResponse({'ok': False, 'error': 'No se recibió imagen.'}, status=400)

    excluir_id = request.POST.get('excluir_id')
    try:
        excluir_id = int(excluir_id) if excluir_id else None
    except (ValueError, TypeError):
        excluir_id = None

    try:
        img_bytes = imagen.read()
        hash_nuevo = calcular_hash_desde_bytes(img_bytes)
    except Exception as e:
        return JsonResponse({'ok': False, 'error': f'Error al procesar imagen: {e}'}, status=422)

    if not hash_nuevo:
        return JsonResponse({'ok': False, 'error': 'No se pudo procesar la imagen.'}, status=422)

    try:
        resultados = buscar_en_todo(hash_nuevo, excluir_marca_id=excluir_id)
    except Exception as e:
        return JsonResponse({'ok': False, 'error': f'Error en la búsqueda: {e}'}, status=500)

    total = len(resultados['marcas']) + len(resultados['predefinidas'])
    es_duplicada = total > 0

    mensaje = None
    if resultados['marcas']:
        mejor = resultados['marcas'][0]
        mensaje = (f'⚠️ Esta imagen ya está registrada para "{mejor["productor"]}" '
                   f'(Marca #{mejor["numero_orden"]}).')
    elif resultados['predefinidas']:
        mejor = resultados['predefinidas'][0]
        mensaje = f'⚠️ Esta imagen ya existe como predefinida: "{mejor["nombre"]}".'

    return JsonResponse({'ok': True, 'es_duplicada': es_duplicada, 'mensaje': mensaje, 'duplicados': resultados})


@login_required
def buscar_marca_por_nombre(request):
    """
    GET /ajax/buscar-marca-nombre/?q=texto
    Busca marcas por nombre de productor, descripción de marca o descripción de señal.
    """
    q = request.GET.get('q', '').strip()
    if len(q) < 2:
        return JsonResponse({'ok': False, 'error': 'Ingresá al menos 2 caracteres.'})

    marcas_qs = MarcaSenal.objects.filter(
        Q(descripcion_marca__icontains=q) |
        Q(descripcion_senal__icontains=q) |
        Q(productor__nombre__icontains=q) |
        Q(productor__apellido__icontains=q) |
        Q(senales_orejeras__icontains=q)
    ).select_related('productor').distinct()[:20]

    marcas = []
    for m in marcas_qs:
        try:
            url = reverse('detalle_marca', kwargs={'pk': m.pk})
        except Exception:
            url = f"/marcas/{m.pk}/"

        orejeras = m.senales_orejeras or {}
        orejeras_txt = []
        if orejeras.get('izquierda'):
            orejeras_txt.append('Izq: ' + ', '.join(orejeras['izquierda']))
        if orejeras.get('derecha'):
            orejeras_txt.append('Der: ' + ', '.join(orejeras['derecha']))

        marcas.append({
            'id': m.pk,
            'numero_orden': m.numero_orden,
            'productor': str(m.productor),
            'productor_id': m.productor.pk,
            'descripcion': m.descripcion_marca[:100],
            'descripcion_senal': (m.descripcion_senal or '')[:100],
            'senales_orejeras_texto': ' | '.join(orejeras_txt),
            'url_detalle': url,
            'imagen_url': m.imagen_marca.url if m.imagen_marca else None,
        })

    total = len(marcas)
    return JsonResponse({'ok': True, 'query': q, 'marcas': marcas, 'total': total})


# ============================================================
# APROBACIÓN DE CUENTAS
# ============================================================

def _es_admin(user):
    try:
        return user.userprofile.rol == 'admin'
    except UserProfile.DoesNotExist:
        return False


@login_required
def cuentas_pendientes(request):
    """Lista de cuentas esperando aprobación (solo admin)."""
    if not _es_admin(request.user):
        messages.error(request, 'No tiene permisos para ver esta sección.')
        return redirect('home')

    pendientes = UserProfile.objects.filter(estado_cuenta='PENDIENTE').select_related('user').order_by('user__date_joined')
    historial = UserProfile.objects.filter(
        estado_cuenta__in=['APROBADA', 'RECHAZADA']
    ).select_related('user', 'aprobado_por').order_by('-fecha_aprobacion')[:100]

    # Recuperar el código generado en su momento (guardado en el ChangeLog)
    codigos_por_profile = {}
    logs_aprobacion = ChangeLog.objects.filter(
        modelo='UserProfile', accion='APROBACION_CUENTA'
    ).order_by('-timestamp')
    for log in logs_aprobacion:
        if log.objeto_id not in codigos_por_profile:
            codigos_por_profile[log.objeto_id] = log.snapshot.get('codigo_generado', '')

    for h in historial:
        h.codigo_historico = codigos_por_profile.get(str(h.pk), '')

    return render(request, 'app_sigrams/cuentas/pendientes.html', {
        'pendientes': pendientes,
        'historial': historial,
    })


@login_required
def aprobar_cuenta(request, user_id):
    """Aprueba una cuenta, le asigna rol y genera un código de activación."""
    if not _es_admin(request.user):
        messages.error(request, 'No tiene permisos para realizar esta acción.')
        return redirect('home')

    usuario = get_object_or_404(User, pk=user_id)
    profile = get_object_or_404(UserProfile, user=usuario)

    if request.method == 'POST':
        rol_elegido = request.POST.get('rol')
        if rol_elegido not in ['empleado', 'inspector', 'admin']:
            messages.error(request, 'Rol inválido.')
            return redirect('cuentas_pendientes')

        import secrets, string
        alfabeto = string.ascii_uppercase + string.digits
        codigo = ''.join(secrets.choice(alfabeto) for _ in range(8))

        profile.rol = rol_elegido
        profile.estado_cuenta = 'APROBADA'
        profile.codigo_activacion = codigo
        profile.aprobado_por = request.user
        profile.fecha_aprobacion = timezone.now()
        profile.save()

        ChangeLog.objects.create(
            user=request.user,
            modelo='UserProfile',
            objeto_id=str(profile.pk),
            accion='APROBACION_CUENTA',
            snapshot={'usuario': usuario.username, 'rol_asignado': rol_elegido, 'codigo_generado': codigo}
        )

        email_enviado = False
        if usuario.email:
            try:
                from django.core.mail import send_mail
                send_mail(
                    subject='SIGRAMS — Tu cuenta fue aprobada',
                    message=(
                        f'Hola {usuario.get_full_name() or usuario.username},\n\n'
                        f'Tu cuenta en SIGRAMS fue aprobada como {profile.get_rol_display()}.\n\n'
                        f'Usuario: {usuario.username}\n'
                        f'Código de activación: {codigo}\n\n'
                        f'Ingresá a la pantalla de activación del sistema con estos datos para activar tu cuenta.'
                    ),
                    from_email=settings.DEFAULT_FROM_EMAIL,
                    recipient_list=[usuario.email],
                    fail_silently=True,
                )
                email_enviado = True
            except Exception:
                email_enviado = False

        return render(request, 'app_sigrams/cuentas/codigo_generado.html', {
            'usuario': usuario,
            'rol_display': profile.get_rol_display(),
            'codigo': codigo,
            'email_enviado': email_enviado,
        })

    return redirect('cuentas_pendientes')


@login_required
def rechazar_cuenta(request, user_id):
    """Rechaza una solicitud de cuenta (queda inactiva permanentemente)."""
    if not _es_admin(request.user):
        messages.error(request, 'No tiene permisos para realizar esta acción.')
        return redirect('home')

    usuario = get_object_or_404(User, pk=user_id)
    profile = get_object_or_404(UserProfile, user=usuario)

    profile.estado_cuenta = 'RECHAZADA'
    profile.aprobado_por = request.user
    profile.fecha_aprobacion = timezone.now()
    profile.save()

    ChangeLog.objects.create(
        user=request.user,
        modelo='UserProfile',
        objeto_id=str(profile.pk),
        accion='RECHAZO_CUENTA',
        snapshot={'usuario': usuario.username}
    )

    messages.success(request, f'Solicitud de cuenta de "{usuario.username}" rechazada.')
    return redirect('cuentas_pendientes')


@ratelimit(key='ip', rate='10/h', method='POST', block=False)
@ratelimit(key='post:username', rate='5/h', method='POST', block=False)
def activar_cuenta_view(request):
    """Pantalla pública donde el usuario ingresa su código para activar la cuenta."""
    if request.method == 'POST':
        print('DEBUG limited =', getattr(request, 'limited', 'SIN ATRIBUTO'))
        # Límite de intentos: por IP y por usuario
        if getattr(request, 'limited', False):
            messages.error(request, 'Demasiados intentos. Esperá un rato antes de volver a probar.')
            return render(request, 'app_sigrams/cuentas/activar.html')

        username = request.POST.get('username', '').strip()
        codigo = request.POST.get('codigo', '').strip().upper()

        # Mismo mensaje para cualquier falla, así no se puede saber si el usuario existe
        mensaje_error = 'Usuario o código de activación incorrecto.'

        usuario = User.objects.filter(username=username).first()
        profile = UserProfile.objects.filter(user=usuario).first() if usuario else None

        if (
            not profile
            or profile.estado_cuenta != 'APROBADA'
            or not profile.codigo_activacion
            or not secrets.compare_digest(codigo, profile.codigo_activacion)
        ):
            messages.error(request, mensaje_error)
            return render(request, 'app_sigrams/cuentas/activar.html')

        usuario.is_active = True
        usuario.save()
        profile.codigo_activacion = None
        profile.save()

        messages.success(request, '¡Cuenta activada! Ya podés iniciar sesión.')
        return redirect('login')

    return render(request, 'app_sigrams/cuentas/activar.html')




# ============================================================
# VISTA PARA ROL INSPECTOR (solo lectura, búsqueda restringida)
# ============================================================

@login_required
def busqueda_inspector(request):
    """
    Pantalla dedicada para el rol Inspector: búsqueda libre por DNI,
    nombre del productor, o descripción de marca/señal. Solo expone
    los campos autorizados (no el resto del sistema).
    """
    q = request.GET.get('q', '').strip()
    resultados = []

    if q:
        marcas_qs = MarcaSenal.objects.filter(
            Q(productor__dni__icontains=q) |
            Q(productor__nombre__icontains=q) |
            Q(productor__apellido__icontains=q) |
            Q(descripcion_marca__icontains=q) |
            Q(descripcion_senal__icontains=q) |
            Q(senales_orejeras__icontains=q)
        ).select_related('productor', 'campo').distinct()[:30]

        for m in marcas_qs:
            resultados.append(m)

    return render(request, 'app_sigrams/inspector/busqueda.html', {
        'q': q,
        'resultados': resultados,
    })


@login_required
def detalle_inspector(request, marca_id):
    """Detalle de solo lectura de una marca, con los campos autorizados para Inspector."""
    marca = get_object_or_404(MarcaSenal.objects.select_related('productor', 'campo'), pk=marca_id)
    return render(request, 'app_sigrams/inspector/detalle.html', {'marca': marca})


# ============================================================
# CARNET FÍSICO (PDF imprimible, frente y dorso)
# ============================================================

@login_required
def generar_carnet_pdf(request, marca_id):
    """Genera el PDF del carnet físico (frente y dorso) para una marca."""
    from reportlab.pdfgen import canvas as pdf_canvas
    from reportlab.lib.units import mm

    marca = get_object_or_404(MarcaSenal.objects.select_related('productor'), pk=marca_id)
    productor = marca.productor

    ANCHO, ALTO = 100 * mm, 70 * mm
    MARGEN = 5 * mm

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="carnet_{marca.numero_orden}.pdf"'

    c = pdf_canvas.Canvas(response, pagesize=(ANCHO, ALTO))
    x0, y0 = MARGEN, MARGEN
    x1, y1 = ANCHO - MARGEN, ALTO - MARGEN
    ancho_util = x1 - x0

    def fit_text(texto, x, y, max_width, font='Helvetica', size=7, min_size=5, centrado=False):
        s = size
        while s > min_size and c.stringWidth(texto, font, s) > max_width:
            s -= 0.25
        c.setFont(font, s)
        if centrado:
            c.drawCentredString(x, y, texto)
        else:
            c.drawString(x, y, texto)

    def hatch_in_circle(cx, cy, r, paso=2.3 * mm):
        path = c.beginPath()
        path.circle(cx, cy, r)
        c.saveState()
        c.clipPath(path, stroke=0, fill=0)
        c.setLineWidth(0.3)
        y = cy - r
        while y <= cy + r:
            c.line(cx - r, y, cx + r, y)
            y += paso
        x = cx - r
        while x <= cx + r:
            c.line(x, cy - r, x, cy + r)
            x += paso
        c.restoreState()
        c.setLineWidth(0.9)
        c.circle(cx, cy, r)

    def hatch_in_oreja(cx, cy, ancho_forma, alto_forma, apunta_izquierda, paso=2.1 * mm):
        """
        Señal de oreja: un extremo vertical y plano, cuerpo ancho con bordes
        superior/inferior curvados, y el otro extremo termina en punta.
        """
        hw = ancho_forma / 2
        hh = alto_forma / 2

        tip_x = cx - hw if apunta_izquierda else cx + hw
        base_x = cx + hw if apunta_izquierda else cx - hw

        path = c.beginPath()

        if apunta_izquierda:
            path.moveTo(tip_x, cy)
            path.curveTo(cx - hw * 0.55, cy + hh * 0.95, cx + hw * 0.35, cy + hh * 0.98, base_x, cy + hh)
            path.lineTo(base_x, cy - hh)
            path.curveTo(cx + hw * 0.35, cy - hh * 0.98, cx - hw * 0.55, cy - hh * 0.95, tip_x, cy)
        else:
            path.moveTo(tip_x, cy)
            path.curveTo(cx + hw * 0.55, cy + hh * 0.95, cx - hw * 0.35, cy + hh * 0.98, base_x, cy + hh)
            path.lineTo(base_x, cy - hh)
            path.curveTo(cx - hw * 0.35, cy - hh * 0.98, cx + hw * 0.55, cy - hh * 0.95, tip_x, cy)

        path.close()

        c.saveState()
        c.clipPath(path, stroke=0, fill=0)
        c.setLineWidth(0.3)
        y = cy - hh
        while y <= cy + hh:
            c.line(cx - hw, y, cx + hw, y)
            y += paso
        x = cx - hw
        while x <= cx + hw:
            c.line(x, cy - hh, x, cy + hh)
            x += paso
        c.restoreState()

        c.setLineWidth(0.9)
        c.drawPath(path, stroke=1, fill=0)

    # ============ FRENTE ============
    c.setLineWidth(1)
    c.rect(x0, y0, ancho_util, y1 - y0)

    cursor = y1 - 3.5 * mm
    c.setFont('Helvetica', 6)
    c.drawString(x0 + 3, cursor, "Gobierno de Catamarca")
    cursor -= 3 * mm
    c.drawString(x0 + 3, cursor, "Ministerio de Producción y Desarrollo")
    cursor -= 3 * mm
    c.drawString(x0 + 3, cursor, "Subsecretaria de Agricultura y Ganaderia")

    c.saveState()
    c.setDash(1, 2)
    caja_w, caja_h = 30 * mm, 9 * mm
    c.rect(x1 - caja_w, y1 - caja_h, caja_w, caja_h)
    c.restoreState()
    c.setFont('Helvetica-Bold', 8)
    c.drawCentredString(x1 - caja_w / 2, y1 - caja_h / 2 - 1.5, f"CARNET N°: {marca.numero_orden}")

    cursor -= 6.5 * mm
    c.setFont('Helvetica-Bold', 12)
    c.drawCentredString(ANCHO / 2, cursor, "Dirección Pcial. de Ganadería")

    cursor -= 4 * mm
    c.line(x0, cursor, x1, cursor)

    cursor -= 5 * mm
    fit_text(f"Apellido y Nombres: {productor.nombre_completo}", x0 + 3, cursor, ancho_util - 6, size=7.5)

    cursor -= 5 * mm
    fit_text(f"L.E./L.C./C.I./D.N.I. N°: {productor.dni}", x0 + 3, cursor, ancho_util - 6, size=7.5)

    cursor -= 3 * mm
    c.line(x0, cursor, x1, cursor)

    domicilio = productor.direccion_completa if hasattr(productor, 'direccion_completa') else (productor.calle or '')
    cursor -= 5 * mm
    fit_text(f"Domicilio Fiscal: {domicilio or '—'}", x0 + 3, cursor, ancho_util - 6, size=7)

    cursor -= 4.5 * mm
    c.setFont('Helvetica', 6.3)
    c.drawString(x0 + 3, cursor, "Registrada el día: ___ / ___ / ___        Registro N°: ______ / ______")

    cursor -= 4.2 * mm
    c.drawString(x0 + 3, cursor, "Renovada el día: ___ / ___ / ___        Registro N°: ______ / ______")

    # Espacio para Sello y Firma
    linea_y = y0 + 8 * mm
    c.setLineWidth(0.6)
    c.line(x0 + 3, linea_y, x0 + 28 * mm, linea_y)
    c.setFont('Helvetica', 6.5)
    c.drawCentredString(x0 + 15.5 * mm, linea_y - 4, "Sello")

    c.line(x1 - 48 * mm, linea_y, x1 - 3, linea_y)
    c.setFont('Helvetica', 5.8)
    c.drawCentredString(x1 - 25.5 * mm, linea_y - 4, "Firma y sello del Enc. Área Marcas y Señales")

    c.showPage()

    # ============ DORSO ============
    c.setLineWidth(1)
    c.rect(x0, y0, ancho_util, y1 - y0)

    cursor = y1 - 4.5 * mm
    vencimiento = marca.fecha_vencimiento.strftime('%d / %m / %Y') if marca.fecha_vencimiento else '___ / ___ / ___'
    c.setFont('Helvetica-Bold', 8)
    c.drawString(x0 + 3, cursor, f"VÁLIDO HASTA EL: {vencimiento}")

    cursor -= 6 * mm
    c.setFont('Helvetica', 6.3)
    c.drawString(x0 + 3, cursor, "Marca para ser aplicada en: _______________     Lado: _______________")

    cursor -= 5 * mm
    c.setFont('Helvetica-Bold', 6.3)
    c.drawString(x0 + 3, cursor, "CARACTERISTICA DE LA SEÑAL:")

    cursor -= 4 * mm
    c.setFont('Helvetica', 6.3)
    c.drawString(x0 + 3, cursor, "O.I.: ___________________________________________")

    cursor -= 4 * mm
    c.drawString(x0 + 3, cursor, "O.D.: ___________________________________________")

    R = 12 * mm
    cx_marca = x0 + 17 * mm
    cy_row = y0 + 18.5 * mm

    hatch_in_circle(cx_marca, cy_row, R)
    c.setFont('Helvetica-Bold', 6.2)
    c.drawCentredString(cx_marca, cy_row - R - 3.5, "MARCA")

    oreja_w, oreja_h = 24 * mm, 13 * mm
    cx_izq = x0 + 46 * mm
    cx_der = x0 + 76 * mm

    hatch_in_oreja(cx_izq, cy_row, oreja_w, oreja_h, apunta_izquierda=True)
    c.setFont('Helvetica-Bold', 6.2)
    c.drawCentredString(cx_izq, cy_row - (oreja_h / 2) - 3.5, "Señal Izquierda")

    hatch_in_oreja(cx_der, cy_row, oreja_w, oreja_h, apunta_izquierda=False)
    c.drawCentredString(cx_der, cy_row - (oreja_h / 2) - 3.5, "Señal Derecha")

    c.showPage()
    c.save()

    if not marca.carnet_generado:
        marca.carnet_generado = True
        marca.fecha_generacion_carnet = timezone.now()
        marca.carnet_generado_por = request.user
        marca.save()

        ChangeLog.objects.create(
            user=request.user,
            modelo='MarcaSenal',
            objeto_id=str(marca.pk),
            accion='GENERACION_CARNET',
            snapshot={'numero_orden': marca.numero_orden}
        )

    return response


@login_required
def historial_eliminaciones(request):
    """Historial de productores eliminados (solo administradores)."""
    if not _es_admin(request.user):
        messages.error(request, 'No tiene permisos para ver esta sección.')
        return redirect('home')

    eliminaciones = ChangeLog.objects.filter(
        modelo='Productor', accion='ELIMINACION'
    ).select_related('user').order_by('-timestamp')[:200]

    return render(request, 'app_sigrams/admin/historial_eliminaciones.html', {
        'eliminaciones': eliminaciones,
    })



